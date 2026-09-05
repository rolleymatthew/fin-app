from __future__ import annotations

import logging
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import List

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Color, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.clients.eastmoney import get_eastmoney_client
from app.clients.eastmoney_datacenter_new import EastmoneyDataNewClient
from app.config import get_settings
from app.constants import spider
from app.constants.cash_flow_type import CashFlowTypeEnum
from app.constants.score import (
    AccountsRecivableScore,
    AssetsScore,
    CashFlowScore,
    GrossProfitScore,
    LiabilScore,
    NetAssetsWeightScore,
    NetProfitScore,
    OperateProfitScore,
    ResultScore,
    StockScore,
)
from app.models.entities import (
    CashFlowScoreEntity,
    FreeCashFlowEntity,
    HKBalanceSheetEntity,
    HKCashFlowEntity,
    HKItemRow,
    HKProfitEntity,
    ScoreEntity,
    YbRoeEntity,
    ZqhFinEntity,
)
from app.repositories.base import MongoRepository
from app.services.kline_service import KLineService
from app.utils import hk_financial_utils, num_utils

logger = logging.getLogger(__name__)

# HK item codes from plan section 8.3
# Balance sheet codes
BS_TOTAL_A = "004009999"
BS_TOTAL_CA = "004002999"
BS_TOTAL_L = "004025999"
BS_TOTAL_CL = "004011999"
BS_TOTAL_NCL = "004020999"
BS_INVENTORY = "004002001"
BS_AR = "004002003"
BS_AP = "004011001"

# Profit codes
PL_REVENUE = "004001001"
PL_TOTAL_REV = "004001999"
PL_COST = "004005001"
PL_OPER_P = "004010999"
PL_PAT_SH = "004025002"
PL_BASIC_EPS = "004027002"
PL_DILUTED_EPS = "004027003"

# Cash flow codes
CF_OPER = "003999"
CF_INVEST = "005999"
CF_FIN = "007999"
CF_PURCH_PPE = "005005"
CF_PURCH_IA = "005007"


class HKFinanceService:
    MAININDICATOR = "RPT_HKF10_FN_MAININDICATOR"
    BALANCE_SHEET = "RPT_HKF10_FN_BALANCE_PC"
    PROFIT = "RPT_HKF10_FN_INCOME_PC"
    CASH_FLOW = "RPT_HKF10_FN_CASHFLOW_PC"

    BALANCE_DATES = "RPT_CUSTOM_HKF10_APPFN_BALANCE_SUMMARY"
    PROFIT_DATES = "RPT_CUSTOM_HKF10_APPFN_INCOME_SUMMARY"
    CASHFLOW_DATES = "RPT_CUSTOM_HKSK_APPFN_CASHFLOW_SUMMARY"

    def __init__(self):
        self.data_new = EastmoneyDataNewClient()
        self.eastmoney = get_eastmoney_client()
        self.kline_service = KLineService()

        self.balance_repo = MongoRepository(HKBalanceSheetEntity)
        self.profit_repo = MongoRepository(HKProfitEntity)
        self.cash_flow_repo = MongoRepository(HKCashFlowEntity)

    def _secucode(self, code: str) -> str:
        return f"{code}.HK"

    # ── Date helpers ──────────────────────────────────────────

    async def _fetch_hk_report_dates(self, code: str, report_name: str) -> List[str]:
        secucode = self._secucode(code)
        dto = await self.data_new.hk_finance(
            report_name,
            "REPORT_DATE",
            f'(SECUCODE="{secucode}")',
            1, 1000, "-1", "REPORT_DATE", "F10", "PC", "07000000000000000",
        )
        results = []
        if dto and dto.get("success"):
            for item in (dto.get("result", {}).get("data") or []):
                rd = str(item.get("REPORT_DATE", "") or "").replace(" 00:00:00", "")
                if rd:
                    results.append(rd)
        return sorted(set(results), reverse=True)

    def _dates_to_filter_batches(self, dates: List[str], batch_size: int = 5) -> List[str]:
        batches = []
        for i in range(0, len(dates), batch_size):
            chunk = dates[i:i + batch_size]
            batches.append(",".join(f"'{d}'" for d in chunk))
        return batches

    async def _fetch_hk_dates_from_data(self, code: str) -> List[str]:
        secucode = self._secucode(code)
        columns = "SECUCODE,SECURITY_CODE,REPORT_DATE"
        dates: set[str] = set()
        for report_name in (self.BALANCE_SHEET, self.PROFIT, self.CASH_FLOW):
            dto = await self.data_new.hk_finance(
                report_name, columns, f'(SECUCODE="{secucode}")',
                1, 1000, "-1", "REPORT_DATE", "HSF10", "PC", "07000000000000000",
            )
            if not dto or not dto.get("success"):
                logger.warning(
                    "[_fetch_hk_dates_from_data] %s secucode=%s 返回空(code=%s msg=%s)",
                    report_name, secucode,
                    dto.get("code") if dto else None,
                    dto.get("message") if dto else None,
                )
                continue
            for item in (dto.get("result", {}).get("data") or []):
                rd = str(item.get("REPORT_DATE", "") or "").replace(" 00:00:00", "")
                if rd:
                    dates.add(rd)
        return sorted(dates, reverse=True)

    # ── Fetch & Save ──────────────────────────────────────────

    async def save_hk_fin_data_to_mongodb(self, code: str, name: str) -> None:
        secucode = self._secucode(code)

        # Fetch and save K-line data
        try:
            kline_entity = await self.kline_service.spider_kline_data(
                code, spider.KLINE_HK_MARKET_CODE, name=name,
            )
            if kline_entity:
                await self.kline_service.save_mongodb(kline_entity)
        except Exception:
            pass

        all_dates = await self._fetch_hk_dates_from_data(code)
        logger.info(
            "[save_hk_fin_data_to_mongodb] code=%s name=%s secucode=%s dates=%s",
            code, name, secucode, all_dates,
        )

        if not all_dates:
            logger.warning(
                "[save_hk_fin_data_to_mongodb] code=%s name=%s 无可用报告日期，跳过写入",
                code, name,
            )
            return

        for batch in self._dates_to_filter_batches(all_dates):
            await self._save_hk_balance_sheet(secucode, code, name, batch)
            await self._save_hk_profit(secucode, code, name, batch)
            await self._save_hk_cash_flow(secucode, code, name, batch)

    async def _save_hk_balance_sheet(self, secucode: str, code: str, name: str, date_filter: str):
        columns = "SECUCODE,SECURITY_CODE,SECURITY_NAME_ABBR,REPORT_DATE,FISCAL_YEAR,STD_ITEM_CODE,STD_ITEM_NAME,AMOUNT"
        filter_str = f'(SECUCODE="{secucode}")(REPORT_DATE in ({date_filter}))'
        dto = await self.data_new.hk_finance(
            self.BALANCE_SHEET, columns, filter_str,
            1, 1000, "-1,1", "REPORT_DATE,STD_ITEM_CODE",
            "F10", "PC", "07000000000000000",
        )
        await self._parse_and_save_row_entity(dto, code, name, self.balance_repo, HKBalanceSheetEntity)

    async def _save_hk_profit(self, secucode: str, code: str, name: str, date_filter: str):
        columns = "SECUCODE,SECURITY_CODE,SECURITY_NAME_ABBR,REPORT_DATE,FISCAL_YEAR,STD_ITEM_CODE,STD_ITEM_NAME,AMOUNT"
        filter_str = f'(SECUCODE="{secucode}")(REPORT_DATE in ({date_filter}))'
        dto = await self.data_new.hk_finance(
            self.PROFIT, columns, filter_str,
            1, 1000, "-1,1", "REPORT_DATE,STD_ITEM_CODE",
            "F10", "PC", "07000000000000000",
        )
        await self._parse_and_save_profit_entity(dto, code, name)

    async def _save_hk_cash_flow(self, secucode: str, code: str, name: str, date_filter: str):
        columns = "SECUCODE,SECURITY_CODE,SECURITY_NAME_ABBR,REPORT_DATE,FISCAL_YEAR,STD_ITEM_CODE,STD_ITEM_NAME,AMOUNT"
        filter_str = f'(SECUCODE="{secucode}")(REPORT_DATE in ({date_filter}))'
        dto = await self.data_new.hk_finance(
            self.CASH_FLOW, columns, filter_str,
            1, 1000, "-1,1", "REPORT_DATE,STD_ITEM_CODE",
            "F10", "PC", "07000000000000000",
        )
        await self._parse_and_save_row_entity(dto, code, name, self.cash_flow_repo, HKCashFlowEntity)

    async def _parse_and_save_row_entity(self, dto, code, name, repo, entity_cls):
        if not dto or not dto.get("success"):
            return
        data = dto.get("result", {}).get("data") or []
        grouped: dict[str, dict] = {}
        for item in data:
            rd = str(item.get("REPORT_DATE", "") or "")
            key = rd
            if key not in grouped:
                grouped[key] = {
                    "securityCode": str(item.get("SECURITY_CODE", code) or code),
                    "securityNameAbbr": str(item.get("SECURITY_NAME_ABBR", name) or name),
                    "reportDate": rd,
                    "fiscalYear": str(item.get("FISCAL_YEAR", "") or ""),
                    "currency": str(item.get("CURRENCY", "") or ""),
                    "items": [],
                }
            grouped[key]["items"].append(HKItemRow(
                itemCode=str(item.get("STD_ITEM_CODE", "") or ""),
                itemName=str(item.get("STD_ITEM_NAME", "") or ""),
                amount=str(item.get("AMOUNT", "") or ""),
            ))
        now = datetime.now().isoformat()
        for rd, doc in grouped.items():
            rd_clean = rd.replace(" 00:00:00", "")
            entity = entity_cls(
                id=f"{code}{rd_clean}",
                securityCode=doc["securityCode"],
                securityNameAbbr=doc["securityNameAbbr"],
                reportDate=rd,
                fiscalYear=doc["fiscalYear"],
                currency=doc.get("currency"),
                items=doc["items"],
                updatedAt=now,
            )
            await repo.save(entity)

    async def _parse_and_save_profit_entity(self, dto, code, name):
        if not dto or not dto.get("success"):
            return
        data = dto.get("result", {}).get("data") or []
        grouped: dict[str, dict] = {}
        for item in data:
            rd = str(item.get("REPORT_DATE", "") or "")
            if rd not in grouped:
                grouped[rd] = {
                    "securityCode": str(item.get("SECURITY_CODE", code) or code),
                    "securityNameAbbr": str(item.get("SECURITY_NAME_ABBR", name) or name),
                    "reportDate": rd,
                    "fiscalYear": str(item.get("FISCAL_YEAR", "") or ""),
                    "basicEps": None,
                    "dilutedEps": None,
                    "items": [],
                }
            item_code = str(item.get("STD_ITEM_CODE", "") or "")
            amount = str(item.get("AMOUNT", "") or "")
            if item_code == PL_BASIC_EPS:
                grouped[rd]["basicEps"] = amount
            elif item_code == PL_DILUTED_EPS:
                grouped[rd]["dilutedEps"] = amount
            grouped[rd]["items"].append(HKItemRow(
                itemCode=item_code,
                itemName=str(item.get("STD_ITEM_NAME", "") or ""),
                amount=amount,
            ))
        now = datetime.now().isoformat()
        for rd, doc in grouped.items():
            rd_clean = rd.replace(" 00:00:00", "")
            entity = HKProfitEntity(
                id=f"{code}{rd_clean}",
                securityCode=doc["securityCode"],
                securityNameAbbr=doc["securityNameAbbr"],
                reportDate=rd,
                fiscalYear=doc["fiscalYear"],
                basicEps=doc["basicEps"],
                dilutedEps=doc["dilutedEps"],
                items=doc["items"],
                updatedAt=now,
            )
            await self.profit_repo.save(entity)

    # ── Main indicator fetch ──────────────────────────────────

    async def _fetch_hk_main_indicator_map(self, code: str) -> dict[str, dict]:
        secucode = self._secucode(code)
        columns = (
            "SECUCODE,SECURITY_CODE,SECURITY_NAME_ABBR,REPORT_DATE,"
            "BASIC_EPS,DILUTED_EPS,EPS_TTM,BPS,"
            "OPERATE_INCOME,OPERATE_INCOME_YOY,"
            "HOLDER_PROFIT,HOLDER_PROFIT_YOY,"
            "GROSS_PROFIT_RATIO,NET_PROFIT_RATIO,ROE_AVG,"
            "DEBT_ASSET_RATIO,CURRENT_RATIO,"
            "ACCOUNTS_RECE_TDAYS,INVENTORY_TDAYS,"
            "TOTAL_ASSETS_TDAYS,CURRENT_ASSETS_TDAYS,"
            "ROIC_YEARLY,TAX_EBT,OCF_SALES,EQUITY_MULTIPLIER"
        )
        filter_str = f'(SECUCODE="{secucode}")'
        dto = await self.data_new.hk_finance(
            self.MAININDICATOR, columns, filter_str,
            1, 1000, "-1", "STD_REPORT_DATE",
            "F10", "PC", "07000000000000000",
        )
        result: dict[str, dict] = {}
        if dto and dto.get("success"):
            for item in (dto.get("result", {}).get("data") or []):
                rd = str(item.get("REPORT_DATE", "") or "").replace(" 00:00:00", "")
                if rd:
                    result[rd] = item
        return result

    # ── Export Excel ──────────────────────────────────────────

    async def export_hk_fin_2_excle(self, code: str, name: str) -> tuple[bool, str]:
        settings = get_settings()
        output_dir = Path(settings.excel_dir).parent / "HKallinone"
        output_dir.mkdir(parents=True, exist_ok=True)

        file_name = output_dir / f"{code}{name}.xlsx"

        balance_list = await self._hk_balance_list(code)
        profit_list = await self._hk_profit_list(code)
        cash_list = await self._hk_cash_flow_list(code)
        main_map = await self._fetch_hk_main_indicator_map(code)

        has_data = len(balance_list) >= 2 and len(profit_list) >= 2 and len(cash_list) >= 2
        if not has_data:
            msg = (
                f"港股财报数据不足：balance={len(balance_list)}, "
                f"profit={len(profit_list)}, cash_flow={len(cash_list)} "
                f"（需各 >= 2 条；请检查 MongoDB 是否写入成功或东方财富接口是否正常）"
            )
            logger.warning("[export_hk_fin_2_excle] %s code=%s name=%s", msg, code, name)
            return False, msg

        base_dir = Path(getattr(sys, "_MEIPASS", Path.cwd()))
        template = base_dir / "app" / "template" / "allin.xlsx"
        if not template.exists():
            template = Path(__file__).resolve().parents[2] / "app" / "template" / "allin.xlsx"
        if not template.exists():
            msg = f"港股财报模板缺失：{template}"
            logger.error("[export_hk_fin_2_excle] %s code=%s name=%s", msg, code, name)
            return False, msg

        wb = load_workbook(template)

        # Sheet 0: 财务透视
        zqh_list = self._build_zqh_list(balance_list, profit_list, cash_list, main_map, code, name)
        if "财务透视" in wb.sheetnames:
            ws0 = wb["财务透视"]
            self._write_zqh_sheet(ws0, zqh_list)

        # Sheet 1: ROE
        yb = await self._build_hk_roe_entity(code, name, main_map)
        if yb is not None and "ROE" in wb.sheetnames:
            ws1 = wb["ROE"]
            self._fill_template(ws1, yb.model_dump() if hasattr(yb, "model_dump") else yb.__dict__)
            await self._write_hk_roe_kline_cells(ws1, code)
        elif "ROE" in wb.sheetnames:
            wb.remove(wb["ROE"])

        # Sheet 2: 现金流状态
        cash_scores = self._build_cash_flow_scores(cash_list, code, name)
        if "现金流状态" in wb.sheetnames:
            ws2 = wb["现金流状态"]
            self._write_cash_flow_sheet(ws2, cash_scores)

        # Sheet 3: 财务评级
        score_entities = self._build_hk_score_entities(balance_list, profit_list, cash_list, main_map, code, name)
        if "财务评级" in wb.sheetnames:
            ws3 = wb["财务评级"]
            self._write_hk_score_sheet(ws3, score_entities)

        # Sheet 4: 股息率分析 (skip for HK - needs separate dividend API)
        # Sheet 5: 自由现金流分析
        fcf_entities = self._build_fcf_entities(cash_list, profit_list, code, name)
        if fcf_entities:
            ws_fcf = wb.create_sheet("自由现金流分析")
            self._write_fcf_sheet(ws_fcf, fcf_entities)

        # Sheet 6+7: 现金转换周期 (skip for HK - needs different calculation)

        wb.save(file_name)
        logger.info(
            "[export_hk_fin_2_excle] 港股财报已落盘 code=%s name=%s path=%s",
            code, name, file_name,
        )
        return True, str(file_name)

    # ── MongoDB queries ──────────────────────────────────────

    async def _hk_balance_list(self, code: str) -> List:
        return await self.balance_repo.find_all_by_security_code_order_by_report_date_desc(code)

    async def _hk_profit_list(self, code: str) -> List:
        return await self.profit_repo.find_all_by_security_code_order_by_report_date_desc(code)

    async def _hk_cash_flow_list(self, code: str) -> List:
        return await self.cash_flow_repo.find_all_by_security_code_order_by_report_date_desc(code)

    # ── Sheet 0: 财务透视 ─────────────────────────────────────

    def _build_zqh_list(self, balance_list, profit_list, cash_list, main_map, code, name):
        profit_by_date = {}
        for p in profit_list:
            rd = (getattr(p, "reportDate", "") or "").replace(" 00:00:00", "")
            if rd:
                profit_by_date[rd] = p
        cash_by_date = {}
        for c in cash_list:
            rd = (getattr(c, "reportDate", "") or "").replace(" 00:00:00", "")
            if rd:
                cash_by_date[rd] = c
        balance_by_date = {}
        for b in balance_list:
            rd = (getattr(b, "reportDate", "") or "").replace(" 00:00:00", "")
            if rd:
                balance_by_date[rd] = b

        all_dates = sorted(set(
            list(profit_by_date.keys()) +
            list(cash_by_date.keys()) +
            list(balance_by_date.keys())
        ), reverse=True)

        result = []
        for rd in all_dates:
            mi = main_map.get(rd, {})
            profit = profit_by_date.get(rd)
            cash = cash_by_date.get(rd)
            balance = balance_by_date.get(rd)

            zqh = ZqhFinEntity()
            zqh.reportDate = rd

            op_inc_raw = mi.get("OPERATE_INCOME")
            if op_inc_raw is not None:
                try:
                    zqh.operatingIncome = round(float(op_inc_raw) / 1e8, 2)
                except (ValueError, TypeError):
                    pass

            oiy = mi.get("OPERATE_INCOME_YOY")
            if oiy is not None:
                try:
                    zqh.revenueGrowthRate = float(oiy)
                except (ValueError, TypeError):
                    pass

            hp_raw = mi.get("HOLDER_PROFIT")
            if hp_raw is not None:
                try:
                    zqh.netProfit = round(float(hp_raw) / 1e8, 2)
                except (ValueError, TypeError):
                    pass

            hpy = mi.get("HOLDER_PROFIT_YOY")
            if hpy is not None:
                try:
                    zqh.netProfitGrowthRate = float(hpy)
                except (ValueError, TypeError):
                    pass

            gpr = mi.get("GROSS_PROFIT_RATIO")
            if gpr is not None:
                try:
                    zqh.operatingGrossProfitMargin = float(gpr)
                except (ValueError, TypeError):
                    pass

            npr = mi.get("NET_PROFIT_RATIO")
            if npr is not None:
                try:
                    zqh.netInterestRate = float(npr)
                except (ValueError, TypeError):
                    pass

            roe_avg = mi.get("ROE_AVG")
            if roe_avg is not None:
                try:
                    zqh.returnOnNetAssets = float(roe_avg)
                except (ValueError, TypeError):
                    pass

            # 营业利润率 = 经营溢利 / 营运收入 × 100
            if profit:
                items = getattr(profit, "items", None)
                oper_p = hk_financial_utils.get_item_amount(items, PL_OPER_P)
                total_rev = hk_financial_utils.get_item_amount(items, PL_TOTAL_REV)
                if oper_p is not None and total_rev is not None and total_rev != 0:
                    zqh.operatingProfitMargin = round(oper_p / total_rev * 100, 2)

            # 经营现金流净额
            if cash:
                items = getattr(cash, "items", None)
                oper_cf = hk_financial_utils.get_item_amount_to_yi(items, CF_OPER)
                if oper_cf is not None:
                    zqh.netOperatingCashFlow = oper_cf

            # 长短期负债比 = 流动负债 / 非流动负债
            if balance:
                items = getattr(balance, "items", None)
                total_cl = hk_financial_utils.get_item_amount(items, BS_TOTAL_CL)
                total_ncl = hk_financial_utils.get_item_amount(items, BS_TOTAL_NCL)
                if total_cl is not None and total_ncl is not None and total_ncl != 0:
                    zqh.lAndLiabRatioww = round(total_cl / total_ncl * 100, 2)

            result.append(zqh)
        return result

    # ── Sheet 1: ROE ─────────────────────────────────────────

    async def _build_hk_roe_entity(self, code: str, name: str, main_map: dict[str, dict]) -> YbRoeEntity | None:
        # Get annual (12-31) report dates
        annual = {k: v for k, v in main_map.items() if k.endswith("-12-31")}
        if len(annual) < 2:
            return None

        years = sorted(annual.keys(), reverse=True)[:4]
        if len(years) < 4:
            return None

        kline_entity = await self.kline_service.kline_by_sec_code(code)
        if not kline_entity or not kline_entity.klines:
            return None

        klines = kline_entity.klines

        year_data = []
        for i, rd in enumerate(years):
            year = int(rd[:4])
            eps_raw = annual[rd].get("BASIC_EPS")
            try:
                eps = Decimal(str(eps_raw)) if eps_raw is not None else None
            except Exception:
                eps = None

            start_date = f"{year}-01-01"
            end_date = f"{year}-12-31"

            year_klines = [
                k for k in klines
                if k.date and start_date <= str(k.date)[:10] <= end_date
            ]
            if not year_klines:
                year_data.append(None)
                continue

            from decimal import ROUND_HALF_UP
            high_v = Decimal(max(num_utils.string_to_double(k.higher) for k in year_klines)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            low_v = Decimal(min(num_utils.string_to_double(k.lower) for k in year_klines)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

            # Average price = total amount / total vol
            total_amount = Decimal(sum(num_utils.string_to_double(k.amount) for k in year_klines if k.amount))
            total_vol = Decimal(sum(num_utils.string_to_double(k.vol) * 100 for k in year_klines if k.vol))
            if total_vol != 0:
                avg_v = (total_amount / total_vol).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            else:
                avg_v = Decimal(0)

            pe_high = (high_v / eps).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) if eps and eps != 0 else Decimal(0)
            pe_avg = (avg_v / eps).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) if eps and eps != 0 else Decimal(0)
            pe_low = (low_v / eps).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) if eps and eps != 0 else Decimal(0)

            year_data.append({
                "date": rd,
                "eps": eps or Decimal(0),
                "high": high_v,
                "avg": avg_v,
                "low": low_v,
                "pe_high": pe_high,
                "pe_avg": pe_avg,
                "pe_low": pe_low,
            })

        if any(x is None for x in year_data):
            return None

        yb = YbRoeEntity()

        for i in range(4):
            d = year_data[i]
            if i == 0:
                yb.lastoneyearhigher = d["high"]
                yb.lastoneyearaverage = d["avg"]
                yb.lastoneyearlower = d["low"]
                yb.oneyeardate = d["date"]
                yb.oneyearepstotal = d["eps"]
                yb.oneyearpedate = d["date"]
                yb.oneyearpehigher = d["pe_high"]
                yb.oneyearpeaverage = d["pe_avg"]
                yb.oneyearpelower = d["pe_low"]
            elif i == 1:
                yb.lasttwoyearhigher = d["high"]
                yb.lasttwoyearaverage = d["avg"]
                yb.lasttwoyearlower = d["low"]
                yb.twoyeardate = d["date"]
                yb.twoyearepstotal = d["eps"]
                yb.twoyearpedate = d["date"]
                yb.twoyearpehigher = d["pe_high"]
                yb.twoyearpeaverage = d["pe_avg"]
                yb.twoyearpelower = d["pe_low"]
            elif i == 2:
                yb.lastthreeyearhigher = d["high"]
                yb.lastthreeyearaverage = d["avg"]
                yb.lastthreeyearlower = d["low"]
                yb.threeyeardate = d["date"]
                yb.threeyearepstotal = d["eps"]
                yb.threeyearpedate = d["date"]
                yb.threeyearpehigher = d["pe_high"]
                yb.threeyearpeaverage = d["pe_avg"]
                yb.threeyearpelower = d["pe_low"]
            elif i == 3:
                yb.lastfouryearhigher = d["high"]
                yb.lastfouryearaverage = d["avg"]
                yb.lastfouryearlower = d["low"]
                yb.fouryeardate = d["date"]
                yb.fouryearepstotal = d["eps"]
                yb.fouryearpedate = d["date"]
                yb.fouryearpehigher = d["pe_high"]
                yb.fouryearpeaverage = d["pe_avg"]
                yb.fouryearpelower = d["pe_low"]

        eps_values = [d["eps"] for d in year_data]
        yb.fouryearavarageepstotal = sum(eps_values, Decimal(0)) / Decimal(4)

        pe_highs = [d["pe_high"] for d in year_data]
        pe_avgs = [d["pe_avg"] for d in year_data]
        pe_lows = [d["pe_low"] for d in year_data]
        yb.fouryearavagpehigher = sum(pe_highs, Decimal(0)) / Decimal(4)
        yb.fouryearavagpemiddle = sum(pe_avgs, Decimal(0)) / Decimal(4)
        yb.fouryearavagpelower = sum(pe_lows, Decimal(0)) / Decimal(4)

        yb.pricehigher = self._calc_price(yb.oneyearepstotal, yb.fouryearavarageepstotal, yb.fouryearavagpehigher)
        yb.pricemiddle = self._calc_price(yb.oneyearepstotal, yb.fouryearavarageepstotal, yb.fouryearavagpemiddle)
        yb.pricelower = self._calc_price(yb.oneyearepstotal, yb.fouryearavarageepstotal, yb.fouryearavagpelower)

        # Fill quarter fields with year data (HK has no quarterly data)
        yb.onequarterdate = yb.oneyeardate
        yb.oneepstotal = yb.oneyearepstotal
        yb.oneepscurrent = yb.oneyearepstotal
        yb.twoquarterdate = yb.twoyeardate
        yb.twoepstotal = yb.twoyearepstotal
        yb.twoepscurrent = yb.twoyearepstotal
        yb.threequarterdate = yb.threeyeardate
        yb.threeepstotal = yb.threeyearepstotal
        yb.threeepscurrent = yb.threeyearepstotal
        yb.fourquarterdate = yb.fouryeardate
        yb.fourepstotal = yb.fouryearepstotal
        yb.fourepscurrent = yb.fouryearepstotal

        yb.onequarterpedate = yb.oneyearpedate
        yb.onequarterpehigher = yb.oneyearpehigher
        yb.onequarterpeaverage = yb.oneyearpeaverage
        yb.onequarterpelower = yb.oneyearpelower
        yb.twoquarterpedate = yb.twoyearpedate
        yb.twoquarterpehigher = yb.twoyearpehigher
        yb.twoquarterpeaverage = yb.twoyearpeaverage
        yb.twoquarterpelower = yb.twoyearpelower
        yb.threequarterpedate = yb.threeyearpedate
        yb.threequarterpehigher = yb.threeyearpehigher
        yb.threequarterpeaverage = yb.threeyearpeaverage
        yb.threequarterpelower = yb.threeyearpelower
        yb.fourquarterpedate = yb.fouryearpedate
        yb.fourquarterpehigher = yb.fouryearpehigher
        yb.fourquarterpeaverage = yb.fouryearpeaverage
        yb.fourquarterpelower = yb.fouryearpelower

        yb.pricequarterhigher = yb.pricehigher
        yb.pricequartermiddle = yb.pricemiddle
        yb.pricequarterlower = yb.pricelower

        yb.fourquarteravagpehigher = yb.fouryearavagpehigher
        yb.fourquarteravagpemiddle = yb.fouryearavagpemiddle
        yb.fourquarteravagpelower = yb.fouryearavagpelower

        yb.laterfirstqepstotal = yb.oneyearepstotal
        yb.latertwoqepstotal = yb.twoyearepstotal
        yb.laterthreeqepstotal = yb.threeyearepstotal
        yb.laterfourqepstotal = yb.fouryearepstotal
        yb.laterfourqavaepstotal = yb.fouryearavarageepstotal

        return yb

    def _calc_price(self, eps: Decimal, avg_eps: Decimal, avg_pe: Decimal) -> Decimal:
        return ((eps + avg_eps) / Decimal(2)) * avg_pe

    async def _write_hk_roe_kline_cells(self, ws, code: str) -> None:
        try:
            kline_entity = await self.kline_service.kline_by_sec_code(code)
        except Exception:
            return
        if not kline_entity or not kline_entity.klines:
            return
        latest = kline_entity.klines[0]
        date_value = getattr(latest, "date", None)
        close_value = getattr(latest, "close", None)
        if date_value is not None:
            ws["J13"].value = date_value
        if close_value is not None:
            ws["J14"].value = close_value

    # ── Sheet 2: 现金流状态 ───────────────────────────────────

    def _build_cash_flow_scores(self, cash_list, code, name):
        result = []
        for c in cash_list:
            if not getattr(c, "reportDate", None):
                continue
            rd = (getattr(c, "reportDate", "") or "").replace(" 00:00:00", "")
            items = getattr(c, "items", None)
            oper = hk_financial_utils.get_item_amount_to_yi(items, CF_OPER)
            invest = hk_financial_utils.get_item_amount_to_yi(items, CF_INVEST)
            fin = hk_financial_utils.get_item_amount_to_yi(items, CF_FIN)

            if oper is None:
                oper = 0.0
            if invest is None:
                invest = 0.0
            if fin is None:
                fin = 0.0

            oper_str = str(round(oper, 2))
            invest_str = str(round(invest, 2))
            fin_str = str(round(fin, 2))

            t = CashFlowTypeEnum.cash_flow_type(oper_str, invest_str, fin_str)

            e = CashFlowScoreEntity()
            e.date = rd
            e.securityCode = code
            e.securityNameAbbr = name
            e.netCashFlowFromOperatingActivities = oper_str
            e.netCashFlowFromInvestmentActivities = invest_str
            e.netCashFlowFromFinancingActivities = fin_str
            e.type = CashFlowTypeEnum.type_name(t)
            e.properties = CashFlowTypeEnum.remark_name(t)
            result.append(e)

        result.sort(key=lambda x: x.date or "", reverse=True)
        return result

    # ── Sheet 3: 财务评级 ─────────────────────────────────────

    def _build_hk_score_entities(self, balance_list, profit_list, cash_list, main_map, code, name):
        profit_by_date = {}
        for p in profit_list:
            rd = (getattr(p, "reportDate", "") or "").replace(" 00:00:00", "")
            profit_by_date[rd] = p
        cash_by_date = {}
        for c in cash_list:
            rd = (getattr(c, "reportDate", "") or "").replace(" 00:00:00", "")
            cash_by_date[rd] = c
        balance_by_date = {}
        for b in balance_list:
            rd = (getattr(b, "reportDate", "") or "").replace(" 00:00:00", "")
            balance_by_date[rd] = b

        all_dates = sorted(set(
            list(profit_by_date.keys()) +
            list(cash_by_date.keys()) +
            list(balance_by_date.keys())
        ), reverse=True)

        result = []
        for rd in all_dates:
            mi = main_map.get(rd, {})
            profit = profit_by_date.get(rd)
            cash = cash_by_date.get(rd)
            balance = balance_by_date.get(rd)

            if not (profit and cash and balance):
                continue

            s = ScoreEntity()
            s.id = f"{code}{rd}"
            s.code = code
            s.name = name
            s.date = rd

            # 毛利率
            gpr = mi.get("GROSS_PROFIT_RATIO")
            try:
                s.grossProfitPer = float(gpr) if gpr is not None else None
            except (ValueError, TypeError):
                s.grossProfitPer = None

            # 营业利润率
            p_items = getattr(profit, "items", None)
            oper_p = hk_financial_utils.get_item_amount(p_items, PL_OPER_P)
            total_rev = hk_financial_utils.get_item_amount(p_items, PL_TOTAL_REV)
            if oper_p is not None and total_rev is not None and total_rev != 0:
                s.operatProfitPer = round(oper_p / total_rev * 100, 2)
            else:
                s.operatProfitPer = None

            # 净利率
            npr = mi.get("NET_PROFIT_RATIO")
            try:
                s.netProfitPer = float(npr) if npr is not None else None
            except (ValueError, TypeError):
                s.netProfitPer = None

            # ROE
            roe_avg = mi.get("ROE_AVG")
            try:
                s.netAssetsWeightPer = float(roe_avg) if roe_avg is not None else None
            except (ValueError, TypeError):
                s.netAssetsWeightPer = None

            # 资产负债率
            b_items = getattr(balance, "items", None)
            total_a = hk_financial_utils.get_item_amount(b_items, BS_TOTAL_A)
            total_l = hk_financial_utils.get_item_amount(b_items, BS_TOTAL_L)
            if total_a is not None and total_l is not None and total_a > 0:
                s.liabilPer = round(total_l / total_a * 100, 4)
            else:
                s.liabilPer = None

            # 流动资产/总资产
            total_ca = hk_financial_utils.get_item_amount(b_items, BS_TOTAL_CA)
            if total_ca is not None and total_a is not None and total_a > 0:
                s.assetsPer = round(total_ca / total_a * 100, 4)
            else:
                s.assetsPer = None

            # 存货周转率 (inventory turnover days from main indicator)
            inv_days = mi.get("INVENTORY_TDAYS")
            try:
                s.stockPer = float(inv_days) if inv_days is not None else None
            except (ValueError, TypeError):
                s.stockPer = None

            # 应收账款周转率 (accounts receivable turnover days)
            ar_days = mi.get("ACCOUNTS_RECE_TDAYS")
            try:
                s.accountsReceivablePer = float(ar_days) if ar_days is not None else None
            except (ValueError, TypeError):
                s.accountsReceivablePer = None

            # 现金流
            c_items = getattr(cash, "items", None)
            oper_cf = hk_financial_utils.get_item_amount(c_items, CF_OPER)
            invest_cf = hk_financial_utils.get_item_amount(c_items, CF_INVEST)
            fin_cf = hk_financial_utils.get_item_amount(c_items, CF_FIN)

            s.netCashFlowFromOperatingActivities = round(oper_cf / 10000, 4) if oper_cf else None
            s.netCashFlowFromInvestmentActivities = round(invest_cf / 10000, 4) if invest_cf else None
            s.netCashFlowFromFinancingActivities = round(fin_cf / 10000, 4) if fin_cf else None

            score_val = (
                GrossProfitScore.includeStart(s.grossProfitPer)
                + OperateProfitScore.includeStart(s.operatProfitPer)
                + NetProfitScore.includeStart(s.netProfitPer)
                + NetAssetsWeightScore.includeStart(s.netAssetsWeightPer)
                + LiabilScore.includeStart(s.liabilPer)
                + AssetsScore.includeStart(s.assetsPer)
                + StockScore.includeStart(s.stockPer)
                + AccountsRecivableScore.includeStart(s.accountsReceivablePer)
                + CashFlowScore.netCashFlowFromOperatingActivities(s.netCashFlowFromOperatingActivities)
                + CashFlowScore.netCashFlowFromInvestmentActivities(s.netCashFlowFromInvestmentActivities)
                + CashFlowScore.netCashFlowFromFinancingActivities(s.netCashFlowFromFinancingActivities)
            )
            s.score = score_val
            s.result = ResultScore.includeStart(score_val)
            s.grossProfitScore = GrossProfitScore.includeStartString(s.grossProfitPer)
            s.operatProfitScore = OperateProfitScore.includeStartString(s.operatProfitPer)
            s.netProfitScore = NetProfitScore.includeStartString(s.netProfitPer)
            s.netAssetsWeightScore = NetAssetsWeightScore.includeStartString(s.netAssetsWeightPer)
            s.liabilScore = LiabilScore.includeStartString(s.liabilPer)
            s.assetsScore = AssetsScore.includeStartString(s.assetsPer)
            s.stockScore = StockScore.includeStartString(s.stockPer)
            s.accountsReceivableScore = AccountsRecivableScore.includeStartString(s.accountsReceivablePer)
            s.netCashFlowFromOperatingActivitiesScore = CashFlowScore.netCashFlowFromOperatingActivitiesString(s.netCashFlowFromOperatingActivities)
            s.netCashFlowFromInvestmentActivitiesScore = CashFlowScore.netCashFlowFromInvestmentActivitiesString(s.netCashFlowFromInvestmentActivities)
            s.netCashFlowFromFinancingActivitiesScore = CashFlowScore.netCashFlowFromFinancingActivitiesString(s.netCashFlowFromFinancingActivities)

            result.append(s)
        return result

    # ── Sheet 5: 自由现金流分析 ─────────────────────────────────

    def _build_fcf_entities(self, cash_list, profit_list, code, name):
        profit_by_date = {}
        for p in profit_list:
            rd = (getattr(p, "reportDate", "") or "").replace(" 00:00:00", "")
            profit_by_date[rd] = p

        result = []
        for c in cash_list:
            rd = (getattr(c, "reportDate", "") or "").replace(" 00:00:00", "")
            if not rd:
                continue
            c_items = getattr(c, "items", None)
            oper_cf = hk_financial_utils.get_item_amount_to_yi(c_items, CF_OPER)
            purch_ppe = hk_financial_utils.get_item_amount(c_items, CF_PURCH_PPE)
            purch_ia = hk_financial_utils.get_item_amount(c_items, CF_PURCH_IA)

            capex = 0.0
            if purch_ppe is not None:
                capex += abs(purch_ppe)
            if purch_ia is not None:
                capex += abs(purch_ia)
            capex_yi = round(capex / 1e8, 2)

            if oper_cf is None:
                continue

            fcf = round(oper_cf - capex_yi, 2)

            e = FreeCashFlowEntity()
            e.date = rd
            e.securityCode = code
            e.securityNameAbbr = name
            e.operatingCashFlow = oper_cf
            e.capex = capex_yi
            e.freeCashFlow = fcf

            profit = profit_by_date.get(rd)
            if profit:
                p_items = getattr(profit, "items", None)
                pat = hk_financial_utils.get_item_amount_to_yi(p_items, PL_PAT_SH)
                if pat and pat > 0:
                    e.fcfToNetProfit = round(fcf / pat * 100, 2)
                else:
                    e.fcfToNetProfit = 0.0
            else:
                e.fcfToNetProfit = 0.0

            result.append(e)

        result.sort(key=lambda x: x.date or "", reverse=True)
        return result

    # ── Excel Writing Helpers (adapted from FinanceService) ───

    def _write_zqh_sheet(self, ws, items: List):
        headers_row1 = [
            "日期/科目",
            "预测公司成长性指标", "预测公司成长性指标", "预测公司成长性指标", "预测公司成长性指标",
            "分析公司获利性指标", "分析公司获利性指标", "分析公司获利性指标", "分析公司获利性指标",
            "检视公司安全性指标", "检视公司安全性指标",
        ]
        headers_row2 = [
            "日期/科目",
            "营业收入(亿元)", "营收增长率%", "净利润(亿元)", "净利润增长率%",
            "营业毛利率%", "净利率%", "营业利润率%", "净资产收益率%",
            "经营现金流净额(亿元)", "长短期负债比",
        ]
        for c, v in enumerate(headers_row1, start=1):
            ws.cell(row=1, column=c, value=v)
        for c, v in enumerate(headers_row2, start=1):
            ws.cell(row=2, column=c, value=v)

        col_widths = [10, 20, 18, 20, 20, 20, 15, 20, 20, 20, 15]
        for i, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 20

        color_grow = Color(indexed=21)
        color_profit = Color(indexed=10)
        color_safe = Color(indexed=24)
        font_grow = Font(name="Microsoft YaHei", size=10, color=color_grow, bold=True)
        font_profit = Font(name="Microsoft YaHei", size=10, color=color_profit, bold=True)
        font_safe = Font(name="Microsoft YaHei", size=10, color=color_safe, bold=True)
        header_fill = PatternFill(fill_type="solid", fgColor="D9D9D9")
        header_align = Alignment(horizontal="center", vertical="center")

        for col in range(1, 12):
            ws.cell(row=1, column=col).fill = header_fill
            ws.cell(row=2, column=col).fill = header_fill
            ws.cell(row=1, column=col).alignment = header_align
            ws.cell(row=2, column=col).alignment = header_align
        for col in range(2, 6):
            ws.cell(row=1, column=col).font = font_grow
            ws.cell(row=2, column=col).font = font_grow
        for col in range(6, 10):
            ws.cell(row=1, column=col).font = font_profit
            ws.cell(row=2, column=col).font = font_profit
        for col in range(10, 12):
            ws.cell(row=1, column=col).font = font_safe
            ws.cell(row=2, column=col).font = font_safe
        ws.cell(row=1, column=1).font = Font(name="Microsoft YaHei", size=10, bold=True)
        ws.cell(row=2, column=1).font = Font(name="Microsoft YaHei", size=10, bold=True)

        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=5)
        ws.merge_cells(start_row=1, start_column=6, end_row=1, end_column=9)
        ws.merge_cells(start_row=1, start_column=10, end_row=1, end_column=11)

        self._write_zqh_data(ws, items, start_row=3)

        if items:
            last_row = 2 + len(items)
            date_align = Alignment(horizontal="center", vertical="center")
            for row in range(3, last_row + 1):
                ws.cell(row=row, column=1).font = Font(name="Calibri", size=11, color="404040", bold=False)
                ws.cell(row=row, column=1).alignment = date_align
                for col in range(2, 6):
                    ws.cell(row=row, column=col).font = Font(name="Calibri", size=11, color=color_grow, bold=False)
                for col in range(6, 10):
                    ws.cell(row=row, column=col).font = Font(name="Calibri", size=11, color=color_profit, bold=False)
                for col in range(10, 12):
                    ws.cell(row=row, column=col).font = Font(name="Calibri", size=11, color=color_safe, bold=False)

    def _write_zqh_data(self, ws, items: List, start_row: int = 1):
        fields = [
            "reportDate", "operatingIncome", "revenueGrowthRate",
            "netProfit", "netProfitGrowthRate",
            "operatingGrossProfitMargin", "netInterestRate",
            "operatingProfitMargin", "returnOnNetAssets",
            "netOperatingCashFlow", "lAndLiabRatioww",
        ]
        for r_idx, item in enumerate(items):
            row = start_row + r_idx
            data = item.model_dump() if hasattr(item, "model_dump") else item.__dict__
            for c_idx, field in enumerate(fields, start=1):
                v = data.get(field)
                if v is not None:
                    cell = ws.cell(row=row, column=c_idx, value=v)
                    if c_idx >= 2 and isinstance(v, (int, float)):
                        cell.number_format = "0.00"

    def _auto_fit_columns(
        self, ws, min_row: int, max_row: int, min_col: int, max_col: int, padding: int = 2
    ):
        for col in range(min_col, max_col + 1):
            max_len = 0
            for row in range(min_row, max_row + 1):
                v = ws.cell(row=row, column=col).value
                if v is None:
                    continue
                text = str(v)
                display_len = 0
                for ch in text:
                    display_len += 2 if ord(ch) > 127 else 1
                if display_len > max_len:
                    max_len = display_len
            ws.column_dimensions[get_column_letter(col)].width = max_len + padding

    def _write_cash_flow_sheet(self, ws, items: List):
        headers = ["日期", "代码", "名称", "经营现金流(亿元)", "投资现金流(亿元)", "筹资现金流(亿元)", "类型", "特征"]
        header_fill = PatternFill(fill_type="solid", fgColor="BFBFBF")
        header_align = Alignment(horizontal="center", vertical="center")
        header_font = Font(name="Microsoft YaHei", size=10, color="00008B", bold=True)
        ws.row_dimensions[1].height = 20
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            ws.column_dimensions[get_column_letter(c)].width = len(str(h)) + 2

        if not items:
            return
        data_font = Font(name="Calibri", size=11, bold=False)
        date_align = Alignment(horizontal="center", vertical="center")
        data_align = Alignment(horizontal="right", vertical="center")

        fields = ["date", "securityCode", "securityNameAbbr",
                   "netCashFlowFromOperatingActivities",
                   "netCashFlowFromInvestmentActivities",
                   "netCashFlowFromFinancingActivities", "type", "properties"]

        type_red = {"妖精型", "老母鸡性", "蛮牛型", "奶牛型"}
        type_green = {"骗吃骗喝型", "混吃等死型", "赌徒型", "大出血型"}
        type_red_fill = PatternFill(fill_type="solid", fgColor="FF0000")
        type_green_fill = PatternFill(fill_type="solid", fgColor="00FF00")
        type_white_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")

        for r_idx, item in enumerate(items):
            row = r_idx + 2
            data = item.model_dump() if hasattr(item, "model_dump") else item.__dict__
            for c_idx, field in enumerate(fields, start=1):
                v = data.get(field)
                cell = ws.cell(row=row, column=c_idx, value=v)
                cell.font = data_font
                cell.alignment = date_align if c_idx == 1 else data_align
                if c_idx in {4, 5, 6} and isinstance(v, (int, float)):
                    cell.number_format = "0.00"
                if c_idx == 7 and v:
                    text = str(v)
                    if text in type_red:
                        cell.font = type_white_font
                        cell.fill = type_red_fill
                    elif text in type_green:
                        cell.font = type_white_font
                        cell.fill = type_green_fill
        self._auto_fit_columns(
            ws, min_row=1, max_row=1 + len(items), min_col=1, max_col=len(headers), padding=2
        )

    def _write_hk_score_sheet(self, ws, items: List):
        headers_row2 = [
            "日期", "总分", "总分评级",
            "毛利率", "毛利率得分",
            "营业利润率", "营业利润率得分",
            "净利率", "净利率得分",
            "净资产收益率", "净资产收益率得分",
            "资产负债比率", "资产负债比率得分",
            "流动资产/总资产", "流动资产/总资产得分",
            "存货周转率", "存货周转率得分",
            "应收账款周转率", "应收账款周转率得分",
            "经营活动现金流量净额", "经营活动现金流量净额得分",
            "投资活动现金流量净额", "投资活动现金流量净额得分",
            "筹资活动现金流量净额", "筹资活动现金流量净额得分",
        ]
        ws.cell(row=1, column=1, value="日期")
        ws.cell(row=1, column=2, value="综合评级")
        ws.cell(row=1, column=4, value="获利能力")
        ws.cell(row=1, column=12, value="财务结构")
        ws.cell(row=1, column=20, value="现金流量")

        header_fill = PatternFill(fill_type="solid", fgColor="BFBFBF")
        header_align = Alignment(horizontal="center", vertical="center")
        header_font = Font(name="Microsoft YaHei", size=10, color="00008B", bold=True)
        profit_font = Font(name="Microsoft YaHei", size=10, color="1F77B4", bold=True)
        structure_font = Font(name="Microsoft YaHei", size=10, color="D62728", bold=True)
        cash_font = Font(name="Microsoft YaHei", size=10, color="2CA02C", bold=True)

        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 20

        for c in range(1, len(headers_row2) + 1):
            ws.cell(row=1, column=c).font = header_font
            ws.cell(row=1, column=c).fill = header_fill
            ws.cell(row=1, column=c).alignment = header_align

        for c, h in enumerate(headers_row2, start=1):
            if c == 1:
                continue
            ws.cell(row=2, column=c, value=h)
            ws.cell(row=2, column=c).font = header_font
            ws.cell(row=2, column=c).fill = header_fill
            ws.cell(row=2, column=c).alignment = header_align
            ws.column_dimensions[get_column_letter(c)].width = len(str(h)) + 2
        ws.column_dimensions[get_column_letter(1)].width = len("日期") + 2

        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
        ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=11)
        ws.merge_cells(start_row=1, start_column=12, end_row=1, end_column=19)
        ws.merge_cells(start_row=1, start_column=20, end_row=1, end_column=25)

        profit_cols = set(range(4, 12))
        structure_cols = set(range(12, 20))
        cash_cols = set(range(20, 26))
        for col in profit_cols:
            ws.cell(row=1, column=col).font = profit_font
            ws.cell(row=2, column=col).font = profit_font
        for col in structure_cols:
            ws.cell(row=1, column=col).font = structure_font
            ws.cell(row=2, column=col).font = structure_font
        for col in cash_cols:
            ws.cell(row=1, column=col).font = cash_font
            ws.cell(row=2, column=col).font = cash_font

        order = [
            "date", "score", "result",
            "grossProfitPer", "grossProfitScore",
            "operatProfitPer", "operatProfitScore",
            "netProfitPer", "netProfitScore",
            "netAssetsWeightPer", "netAssetsWeightScore",
            "liabilPer", "liabilScore",
            "assetsPer", "assetsScore",
            "stockPer", "stockScore",
            "accountsReceivablePer", "accountsReceivableScore",
            "netCashFlowFromOperatingActivities", "netCashFlowFromOperatingActivitiesScore",
            "netCashFlowFromInvestmentActivities", "netCashFlowFromInvestmentActivitiesScore",
            "netCashFlowFromFinancingActivities", "netCashFlowFromFinancingActivitiesScore",
        ]
        if not items:
            return
        data_font = Font(name="Calibri", size=11, bold=False)
        data_profit_font = Font(name="Calibri", size=11, bold=False, color="1F77B4")
        data_structure_font = Font(name="Calibri", size=11, bold=False, color="D62728")
        data_cash_font = Font(name="Calibri", size=11, bold=False, color="2CA02C")
        data_align = Alignment(horizontal="right", vertical="center")
        date_align = Alignment(horizontal="center", vertical="center")
        rating_good_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        rating_mid_font = Font(name="Microsoft YaHei", size=11, bold=True, color="000000")
        rating_watch_font = Font(name="Microsoft YaHei", size=11, bold=True, color="006400")
        rating_good_fill = PatternFill(fill_type="solid", fgColor="FF0000")
        rating_mid_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
        rating_watch_fill = PatternFill(fill_type="solid", fgColor="00FF00")

        profit_cols_data = set(range(4, 12))
        structure_cols_data = set(range(12, 20))
        cash_cols_data = set(range(20, 26))

        for r_idx, item in enumerate(items):
            row = r_idx + 3
            data = item.model_dump() if hasattr(item, "model_dump") else item.__dict__
            for c_idx, key in enumerate(order, start=1):
                v = data.get(key)
                if key == "date" and v and isinstance(v, str) and v.endswith(" 00:00:00"):
                    v = v.replace(" 00:00:00", "")
                cell = ws.cell(row=row, column=c_idx, value=v)
                if v is not None:
                    if c_idx == 3:
                        rating_text = str(v)
                        if rating_text == "优等":
                            cell.font = rating_good_font
                            cell.fill = rating_good_fill
                        elif rating_text == "中等":
                            cell.font = rating_mid_font
                            cell.fill = rating_mid_fill
                        elif rating_text == "观望":
                            cell.font = rating_watch_font
                            cell.fill = rating_watch_fill
                        else:
                            cell.font = data_font
                    elif c_idx in profit_cols_data:
                        cell.font = data_profit_font
                    elif c_idx in structure_cols_data:
                        cell.font = data_structure_font
                    elif c_idx in cash_cols_data:
                        cell.font = data_cash_font
                    else:
                        cell.font = data_font
                    if c_idx not in {1, 3} and isinstance(v, (int, float)):
                        cell.number_format = "0.00"
                    cell.alignment = date_align if c_idx in {1, 3} else data_align
        self._auto_fit_columns(
            ws, min_row=2, max_row=2 + len(items), min_col=1, max_col=len(headers_row2), padding=2
        )

    def _write_fcf_sheet(self, ws, items: List):
        headers = ["日期", "代码", "名称", "经营现金流(亿元)", "资本开支(亿元)", "自由现金流(亿元)", "FCF/净利润(%)"]
        header_fill = PatternFill(fill_type="solid", fgColor="BFBFBF")
        header_align = Alignment(horizontal="center", vertical="center")
        header_font = Font(name="Microsoft YaHei", size=10, color="00008B", bold=True)
        ws.row_dimensions[1].height = 20
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            ws.column_dimensions[get_column_letter(c)].width = len(str(h)) + 2

        if not items:
            return
        order = ["date", "securityCode", "securityNameAbbr", "operatingCashFlow", "capex", "freeCashFlow", "fcfToNetProfit"]
        data_font = Font(name="Calibri", size=11, bold=False)
        data_align = Alignment(horizontal="right", vertical="center")
        date_align = Alignment(horizontal="center", vertical="center")
        for r_idx, item in enumerate(items):
            row = r_idx + 2
            data = item.model_dump() if hasattr(item, "model_dump") else item.__dict__
            for c_idx, field in enumerate(order, start=1):
                v = data.get(field)
                cell = ws.cell(row=row, column=c_idx, value=v)
                cell.font = data_font
                cell.alignment = date_align if c_idx == 1 else data_align
                if c_idx in {4, 5, 6, 7} and isinstance(v, (int, float)):
                    cell.number_format = "0.00"
        self._auto_fit_columns(
            ws, min_row=1, max_row=1 + len(items), min_col=1, max_col=len(headers), padding=2
        )

    def _fill_template(self, ws, data: dict):
        if not data:
            return
        safe = {str(k): "" if v is None else str(v) for k, v in data.items()}
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str) or "{" not in v or "}" not in v:
                    continue
                new_v = v
                for k, val in safe.items():
                    token = "{" + k + "}"
                    if token in new_v:
                        new_v = new_v.replace(token, val)
                if new_v != v:
                    cell.value = new_v
