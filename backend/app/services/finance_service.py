from __future__ import annotations

import logging
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import List

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.clients.eastmoney_datacenter import EastmoneyDataCenterClient
from app.clients.eastmoney_datacenter_new import EastmoneyDataNewClient
from app.clients.finance_eastmoney import FinanceEastmoneyClient
from app.clients.finance_eastmoney_v2 import FinanceEastmoneyV2Client
from app.config import get_settings
from app.constants.org_type import (
    BankTypeCode,
    InsuranceTypeCode,
    SecuritiesTypeCode,
    UniversalTypeCode,
)
from app.constants.score import (
    AccountsRecivableScore,
    AssetsScore,
    CashFlowScore,
    GrossProfitBankScore,
    GrossProfitScore,
    LiabilScore,
    NetAssetsWeightBankScore,
    NetAssetsWeightScore,
    NetProfitBankScore,
    NetProfitScore,
    OperateProfitBankScore,
    OperateProfitScore,
    ResultScore,
    StockScore,
)
from app.mappers.cash_score import to_entity as cash_score_to_entity
from app.mappers.common import dto_to_entity
from app.mappers.custom import (
    bonus_dto_to_entity,
    dupond_dto_to_entity,
    fin_to_entity,
    share_bonus_dto_to_entity,
)
from app.mappers.yb_eps import YbEpsEntityMapper
from app.mappers.zqh import (
    creat_for_universal,
    creats_for_bank,
    creats_for_bond,
    creats_for_insurance,
)
from app.models.dto import RoeDTO
from app.models.entities import (
    AssetsBankEntity,
    AssetsInsuranceEntity,
    AssetsSecuritiesEntity,
    AssetsUniversalEntity,
    BonusEntity,
    CashFlowBankEntity,
    CashFlowInsuranceEntity,
    CashFlowSecuritiesEntity,
    CashFlowUniversalEntity,
    DividendYieldEntity,
    DuPondEntity,
    FinEntity,
    FreeCashFlowEntity,
    ProfitBankEntity,
    ProfitForecastEntity,
    ProfitInsuranceEntity,
    ProfitSecuritiesEntity,
    ProfitUniversalEntity,
    ScoreEntity,
    SecCodeEntity,
    ShareBonusEntity,
    YbRoeEntity,
    YbRoeEpsEntity,
    ZqhFinEntity,
)
from app.repositories.base import MongoRepository
from app.services.kline_service import KLineService
from app.services.seccode_service import SecCodeService
from app.services.cash_cycle_service import CashCycleService
from app.utils import date_utils, finance_utils, num_utils, transform

logger = logging.getLogger(__name__)


class FinanceService:
    PROFIT = "profit"
    CASH_FLOW = "cashflow"
    ASSETS = "assets"

    def __init__(self):
        if get_settings().use_finance_eastmoney_v2:
            self.finance_client = FinanceEastmoneyV2Client()
        else:
            self.finance_client = FinanceEastmoneyClient()
        self.data_center = EastmoneyDataCenterClient()
        self.data_new = EastmoneyDataNewClient()

        self.sec_code_service = SecCodeService()
        self.kline_service = KLineService()

        self.profit_universal_repo = MongoRepository(ProfitUniversalEntity)
        self.profit_securities_repo = MongoRepository(ProfitSecuritiesEntity)
        self.profit_insurance_repo = MongoRepository(ProfitInsuranceEntity)
        self.profit_bank_repo = MongoRepository(ProfitBankEntity)

        self.cash_universal_repo = MongoRepository(CashFlowUniversalEntity)
        self.cash_securities_repo = MongoRepository(CashFlowSecuritiesEntity)
        self.cash_insurance_repo = MongoRepository(CashFlowInsuranceEntity)
        self.cash_bank_repo = MongoRepository(CashFlowBankEntity)

        self.assets_universal_repo = MongoRepository(AssetsUniversalEntity)
        self.assets_securities_repo = MongoRepository(AssetsSecuritiesEntity)
        self.assets_insurance_repo = MongoRepository(AssetsInsuranceEntity)
        self.assets_bank_repo = MongoRepository(AssetsBankEntity)

        self.dupond_repo = MongoRepository(DuPondEntity)
        self.bonus_repo = MongoRepository(BonusEntity)
        self.share_bonus_repo = MongoRepository(ShareBonusEntity)
        self.fin_repo = MongoRepository(FinEntity)
        self.score_repo = MongoRepository(ScoreEntity)
        self.roe_eps_repo = MongoRepository(YbRoeEpsEntity)

    async def profit_universal_list_by_code(self, sec_code: str):
        return await self.profit_universal_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def profit_bank_list_by_code(self, sec_code: str):
        return await self.profit_bank_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def profit_securities_list_by_code(self, sec_code: str):
        return await self.profit_securities_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def profit_insurance_list_by_code(self, sec_code: str):
        return await self.profit_insurance_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def cash_flow_universal_list_by_code(self, sec_code: str):
        return await self.cash_universal_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def cash_flow_bank_list_by_code(self, sec_code: str):
        return await self.cash_bank_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def cash_flow_securities_list_by_code(self, sec_code: str):
        return await self.cash_securities_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def cash_flow_insurance_list_by_code(self, sec_code: str):
        return await self.cash_insurance_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def assets_universal_list_by_code(self, sec_code: str):
        return await self.assets_universal_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def assets_bank_list_by_code(self, sec_code: str):
        return await self.assets_bank_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def assets_securities_list_by_code(self, sec_code: str):
        return await self.assets_securities_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def assets_insurance_list_by_code(self, sec_code: str):
        return await self.assets_insurance_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )

    async def roe_list_by_code(self, sec_code: str):
        dupond = await self.dupond_repo.find_all_by_security_code_order_by_report_date_desc(
            sec_code
        )
        items = []
        for d in dupond:
            data = d.model_dump() if hasattr(d, "model_dump") else d.__dict__
            extra = getattr(d, "__pydantic_extra__", None)
            if isinstance(extra, dict) and extra:
                data = {**data, **extra}
            items.append(RoeDTO.model_validate(data))
        return items

    async def share_bonus_list_by_code(self, sec_code: str):
        return await self.share_bonus_repo.find_all_by_security_code_order_by_report_date_desc(sec_code)

    async def _get_profit_list_by_org_type(self, sec_code_entity: SecCodeEntity):
        sec_code = sec_code_entity.securityCode
        org_type = sec_code_entity.orgTypeCode
        if org_type == UniversalTypeCode:
            return await self.profit_universal_list_by_code(sec_code)
        elif org_type == BankTypeCode:
            return await self.profit_bank_list_by_code(sec_code)
        elif org_type == InsuranceTypeCode:
            return await self.profit_insurance_list_by_code(sec_code)
        elif org_type == SecuritiesTypeCode:
            return await self.profit_securities_list_by_code(sec_code)
        return []

    async def get_sec_code_entities(self, sec_codes: List[str] | None):
        entities = []
        if sec_codes:
            for s in sec_codes:
                e = await self.sec_code_service.sec_code_entity_by_id(s)
                if e:
                    entities.append(e)
        else:
            entities = await self.sec_code_service.sec_code_list()
        return [s for s in entities if finance_utils.null_st(s)]

    async def get_sec_code_entities_with_crawl(self, sec_codes: List[str] | None):
        if sec_codes:
            for s in sec_codes:
                await self.sec_code_service.sec_code_entity_by_one(s)
        return await self.get_sec_code_entities(sec_codes)

    def get_sec_code_split_lists(self, sec_codes: List[str], keep_on_code: str | None):
        collect = sorted([x for x in sec_codes if x])
        if keep_on_code:
            collect = [x for x in collect if x > keep_on_code]
        part = 10
        return [collect[i : i + part] for i in range(0, len(collect), part)]

    async def save_fin_data_to_mongodb(
        self, sec_code_entity: SecCodeEntity, date: str | None, is_crawl: bool | None
    ):
        if sec_code_entity is None:
            return None
        try:
            split = sec_code_entity.secucode.split(".") if sec_code_entity.secucode else []
            if len(split) == 2:
                sh_code = f"{split[1]}{split[0]}"
                c_type = sec_code_entity.orgTypeCode
                if is_crawl:
                    await self.save_du_pont_data(sec_code_entity)
                    await self.save_share_bonus_all_data(sec_code_entity)
                if c_type == SecuritiesTypeCode:
                    return await self._save_securities_data(date, sec_code_entity, sh_code, c_type)
                if c_type == InsuranceTypeCode:
                    return await self._save_insurance_data(date, sec_code_entity, sh_code, c_type)
                if c_type == BankTypeCode:
                    return await self._save_bank_data(date, sec_code_entity, sh_code, c_type)
                if c_type == UniversalTypeCode:
                    return await self._save_universal_data(date, sec_code_entity, sh_code, c_type)
        except Exception:
            logger.exception(
                "[save_fin_data_to_mongodb] failed sec_code=%s secucode=%s",
                sec_code_entity.securityCode,
                sec_code_entity.secucode,
            )
            return sec_code_entity.securityCode
        return None

    async def _save_bank_data(
        self, date: str | None, sec_code_entity: SecCodeEntity, sh_code: str, c_type: str
    ):
        date_map = await self._get_date_map(date, sh_code, c_type)
        if not (self.PROFIT in date_map and self.CASH_FLOW in date_map and self.ASSETS in date_map):
            return sec_code_entity.securityCode

        profit_list = []
        for d in date_map[self.PROFIT]:
            text = await self.finance_client.profit(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                profit_list.append(dto_to_entity(item, ProfitBankEntity))
        cash_list = []
        for d in date_map[self.CASH_FLOW]:
            text = await self.finance_client.cash_flow(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                cash_list.append(dto_to_entity(item, CashFlowBankEntity))
        assets_list = []
        for d in date_map[self.ASSETS]:
            text = await self.finance_client.assets(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                assets_list.append(dto_to_entity(item, AssetsBankEntity))

        await self.profit_bank_repo.save_many(profit_list)
        await self.cash_bank_repo.save_many(cash_list)
        await self.assets_bank_repo.save_many(assets_list)
        return None

    async def _save_insurance_data(
        self, date: str | None, sec_code_entity: SecCodeEntity, sh_code: str, c_type: str
    ):
        date_map = await self._get_date_map(date, sh_code, c_type)
        if not (self.PROFIT in date_map and self.CASH_FLOW in date_map and self.ASSETS in date_map):
            return sec_code_entity.securityCode

        profit_list = []
        for d in date_map[self.PROFIT]:
            text = await self.finance_client.profit(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                profit_list.append(dto_to_entity(item, ProfitInsuranceEntity))
        cash_list = []
        for d in date_map[self.CASH_FLOW]:
            text = await self.finance_client.cash_flow(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                cash_list.append(dto_to_entity(item, CashFlowInsuranceEntity))
        assets_list = []
        for d in date_map[self.ASSETS]:
            text = await self.finance_client.assets(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                assets_list.append(dto_to_entity(item, AssetsInsuranceEntity))

        await self.profit_insurance_repo.save_many(profit_list)
        await self.cash_insurance_repo.save_many(cash_list)
        await self.assets_insurance_repo.save_many(assets_list)
        return None

    async def _save_securities_data(
        self, date: str | None, sec_code_entity: SecCodeEntity, sh_code: str, c_type: str
    ):
        date_map = await self._get_date_map(date, sh_code, c_type)
        if not (self.PROFIT in date_map and self.CASH_FLOW in date_map and self.ASSETS in date_map):
            return sec_code_entity.securityCode

        profit_list = []
        for d in date_map[self.PROFIT]:
            text = await self.finance_client.profit(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                profit_list.append(dto_to_entity(item, ProfitSecuritiesEntity))
        cash_list = []
        for d in date_map[self.CASH_FLOW]:
            text = await self.finance_client.cash_flow(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                cash_list.append(dto_to_entity(item, CashFlowSecuritiesEntity))
        assets_list = []
        for d in date_map[self.ASSETS]:
            text = await self.finance_client.assets(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                assets_list.append(dto_to_entity(item, AssetsSecuritiesEntity))

        await self.profit_securities_repo.save_many(profit_list)
        await self.cash_securities_repo.save_many(cash_list)
        await self.assets_securities_repo.save_many(assets_list)
        return None

    async def _save_universal_data(
        self, date: str | None, sec_code_entity: SecCodeEntity, sh_code: str, c_type: str
    ):
        date_map = await self._get_date_map(date, sh_code, c_type)
        if not (self.PROFIT in date_map and self.CASH_FLOW in date_map and self.ASSETS in date_map):
            return sec_code_entity.securityCode

        profit_list = []
        for d in date_map[self.PROFIT]:
            text = await self.finance_client.profit(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                profit_list.append(dto_to_entity(item, ProfitUniversalEntity))
        cash_list = []
        for d in date_map[self.CASH_FLOW]:
            text = await self.finance_client.cash_flow(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                cash_list.append(dto_to_entity(item, CashFlowUniversalEntity))
        assets_list = []
        for d in date_map[self.ASSETS]:
            text = await self.finance_client.assets(c_type, d, sh_code)
            dto = transform.normalize_keys(__import__("json").loads(text)) if text else {}
            for item in dto.get("data", []) or []:
                assets_list.append(dto_to_entity(item, AssetsUniversalEntity))

        await self.profit_universal_repo.save_many(profit_list)
        await self.cash_universal_repo.save_many(cash_list)
        await self.assets_universal_repo.save_many(assets_list)
        return None

    async def _get_date_map(self, date: str | None, sh_code: str, c_type: str):
        profit_dates = []
        cash_dates = []
        assets_dates = []
        if not date:
            profit_dto = await self.finance_client.profit_dates(c_type, sh_code)
            if profit_dto and profit_dto.get("data"):
                profit_dates.extend(self._get_date_list(profit_dto))
            cash_dto = await self.finance_client.cash_flow_dates(c_type, sh_code)
            if cash_dto and cash_dto.get("data"):
                cash_dates.extend(self._get_date_list(cash_dto))
            assets_dto = await self.finance_client.assets_dates(c_type, sh_code)
            if assets_dto and assets_dto.get("data"):
                assets_dates.extend(self._get_date_list(assets_dto))
        else:
            profit_dates.append(date)
            cash_dates.append(date)
            assets_dates.append(date)
        return {self.PROFIT: profit_dates, self.CASH_FLOW: cash_dates, self.ASSETS: assets_dates}

    def _get_date_list(self, date_dto: dict) -> List[str]:
        date_list = []
        size = 5
        data = date_dto.get("data", [])

        def _get_report_date(item: dict) -> str:
            if not item:
                return ""
            v = item.get("reportDate")
            if not v:
                v = item.get("REPORT_DATE") or item.get("REPORTDATE")
            return str(v or "")

        dates = [_get_report_date(d).replace(" 00:00:00", "") for d in data if _get_report_date(d)]
        # Ensure latest dates are fetched first regardless of upstream ordering
        dates = sorted(set(dates), reverse=True)
        if not dates:
            return []
        page = len(dates) // size
        if len(dates) > size * page:
            page += 1
        for i in range(page):
            skip = i * size
            collect1 = ",".join(dates[skip : skip + size])
            # Let HTTP client handle URL-encoding; avoid double-encoding here
            date_list.append(collect1)
        return date_list

    async def save_du_pont_data(self, sec_code_entity: SecCodeEntity):
        columns_name = (
            "SECUCODE,SECURITY_CODE,SECURITY_NAME_ABBR,ORG_CODE,ORG_TYPE,REPORT_DATE,REPORT_TYPE,REPORT_DATE_NAME,SECURITY_TYPE_CODE,NOTICE_DATE,UPDATE_DATE,CURRENCY,"
            "NETPROFIT,TOTAL_OPERATE_INCOME,TOTAL_ASSETS,TOTAL_LIABILITIES,TOTAL_CURRENT_ASSETS,TOTAL_NONCURRENT_ASSETS,PARENT_NETPROFIT,SALE_NPR,TOTAL_ASSETS_TR,JROA,PARENT_NETPROFIT_RATIO,"
            "EQUITY_MULTIPLIER,ROE,DEBT_ASSET_RATIO,TOTAL_INCOME,TOTAL_COST,TOTAL_EXPENSE,MONETARYFUNDS,TRADE_FINASSET,NOTE_RECE,ACCOUNTS_RECE,FINANCE_RECE,OTHER_RECE,INVENTORY,CREDITOR_INVEST,LONG_EQUITY_INVEST,"
            "INVEST_REALESTATE,FIXED_ASSET,CIP,USERIGHT_ASSET,INTANGIBLE_ASSET,DEVELOP_EXPENSE,GOODWILL,LONG_PREPAID_EXPENSE,DEFER_TAX_ASSET,INVEST_INCOME,EXCHANGE_INCOME,FAIRVALUE_CHANGE_INCOME,ASSET_DISPOSAL_INCOME,"
            "OPERATE_COST,SURRENDER_VALUE,NET_COMPENSATE_EXPENSE,NET_CONTRACT_RESERVE,POLICY_BONUS_EXPENSE,OPERATE_TAX_ADD,INCOME_TAX,ASSET_IMPAIRMENT_INCOME,CREDIT_IMPAIRMENT_INCOME,NONBUSINESS_EXPENSE,FINANCE_EXPENSE,SALE_EXPENSE,MANAGE_EXPENSE,RESEARCH_EXPENSE,"
            "INTEREST_NI,FEE_COMMISSION_NI,EARNED_PREMIUM,BUSINESS_MANAGE_EXPENSE,OTHER_CREDITOR_INVEST,OTHER_EQUITY_INVEST,LONG_RECE,AVAILABLE_SALE_FINASSET,HOLD_MATURITY_INVEST,FEE_COMMISSION_EXPENSE"
        )
        columns = columns_name
        filter_str = f'(SECUCODE="{sec_code_entity.secucode}")'
        dto = await self.data_new.dupond(
            "RPT_F10_FINANCE_DUPONT",
            columns,
            filter_str,
            1,
            1000,
            "-1",
            "REPORT_DATE",
            "HSF10",
            "PC",
            "024920210926389585",
        )
        if not dto or not dto.get("success"):
            print(f"[dupond] empty secucode={sec_code_entity.secucode} dto={dto}", flush=True)
            return
        data_list = dto.get("result", {}).get("data", []) or []
        for item in data_list:
            entity = dupond_dto_to_entity(item)
            await self.dupond_repo.save(entity)

    async def save_share_bonus_all_data(self, sec_code_entity: SecCodeEntity):
        filter_str = f'(SECUCODE="{sec_code_entity.secucode}")'
        bonus_dto = await self.data_new.bonus(
            "RPT_F10_DIVIDEND_MAIN",
            "ALL",
            filter_str,
            1,
            100,
            "-1",
            "NOTICE_DATE",
            "HSF10",
            "PC",
            "0913818412108917",
        )
        share_bonus_dto = await self.data_center.share_bonus(
            "REPORT_DATE", -1, 300, 1, "RPT_SHAREBONUS_DET", "ALL", "WEB", "WEB", filter_str
        )
        if (
            not bonus_dto
            or not bonus_dto.get("success")
            or not share_bonus_dto
            or not share_bonus_dto.get("success")
        ):
            return
        for item in bonus_dto.get("result", {}).get("data", []) or []:
            await self.bonus_repo.save(bonus_dto_to_entity(item))
        for item in share_bonus_dto.get("result", {}).get("data", []) or []:
            await self.share_bonus_repo.save(share_bonus_dto_to_entity(item))

    async def export_fin_to_excel(self, sec_code_entity: SecCodeEntity, force: int):
        zqh_list = await self.proceed_zqh(sec_code_entity)
        cash_items = await self.cash_flow_universal_list_by_code(sec_code_entity.securityCode)
        cash_scores = [cash_score_to_entity(c) for c in cash_items]
        score_entities = await self.get_score_entities(sec_code_entity)
        await self.score_repo.save_many(score_entities)

        yb = await self._get_yb_roe_entity(force, sec_code_entity)
        if yb is not None:
            await self._save_roe_mongo(yb, sec_code_entity)

        settings = get_settings()

        # Prefer bundled resource when running as exe; fallback to cwd/repo layout.
        base_dir = Path(getattr(sys, "_MEIPASS", Path.cwd()))
        template = base_dir / "app" / "template" / "allin.xlsx"
        if not template.exists():
            template = Path(__file__).resolve().parents[2] / "app" / "template" / "allin.xlsx"
        if not template.exists():
            return

        output_dir = Path(settings.excel_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        file_name = (
            output_dir / f"py{sec_code_entity.securityCode}{sec_code_entity.securityNameAbbr}.xlsx"
        )

        wb = load_workbook(template)
        ws0 = wb["财务透视"]
        ws2 = wb["现金流状态"]
        ws3 = wb["财务评级"]

        self._write_zqh_sheet(ws0, zqh_list)
        if yb is not None:
            ws1 = wb["ROE"]
            self._fill_template(ws1, yb.model_dump() if hasattr(yb, "model_dump") else yb.__dict__)
            await self._write_roe_kline_cells(ws1, sec_code_entity)
        else:
            if "ROE" in wb.sheetnames:
                wb.remove(wb["ROE"])
        if sec_code_entity.orgTypeCode not in {BankTypeCode, InsuranceTypeCode, SecuritiesTypeCode}:
            self._write_cash_flow_sheet(ws2, cash_scores)
        else:
            if "现金流状态" in wb.sheetnames:
                wb.remove(wb["现金流状态"])
        self._write_score_sheet(ws3, score_entities, sec_code_entity.orgTypeCode)

        dividend_entities = await self._calc_dividend_yield_entities(sec_code_entity)
        ws4 = wb.create_sheet("股息率分析")
        self._write_dividend_sheet(ws4, dividend_entities)

        if sec_code_entity.orgTypeCode == UniversalTypeCode:
            ccc_service = CashCycleService()
            ccc_records = await ccc_service.calc_cash_cycle(sec_code_entity.securityCode)
            if ccc_records:
                ws5 = wb.create_sheet("现金转换周期")
                CashCycleService.write_to_sheet(ws5, ccc_records)
            annual_records = await ccc_service.calc_annual_ccc(sec_code_entity.securityCode)
            if annual_records:
                ws6 = wb.create_sheet("年度CCC")
                CashCycleService.write_annual_sheet(ws6, annual_records)
            fcf_entities = await self._build_fcf_entities(sec_code_entity)
            if fcf_entities:
                ws7 = wb.create_sheet("自由现金流分析")
                self._write_fcf_sheet(ws7, fcf_entities)

        wb.save(file_name)

    async def proceed_zqh(self, sec_code_entity: SecCodeEntity) -> List[ZqhFinEntity]:
        if sec_code_entity.orgTypeCode == UniversalTypeCode:
            zqh = creat_for_universal(
                await self.profit_universal_list_by_code(sec_code_entity.securityCode),
                await self.cash_flow_universal_list_by_code(sec_code_entity.securityCode),
                await self.assets_universal_list_by_code(sec_code_entity.securityCode),
                await self.roe_list_by_code(sec_code_entity.securityCode),
                sec_code_entity,
            )
        elif sec_code_entity.orgTypeCode == BankTypeCode:
            zqh = creats_for_bank(
                await self.profit_bank_list_by_code(sec_code_entity.securityCode),
                await self.cash_flow_bank_list_by_code(sec_code_entity.securityCode),
                await self.roe_list_by_code(sec_code_entity.securityCode),
                sec_code_entity,
            )
        elif sec_code_entity.orgTypeCode == InsuranceTypeCode:
            zqh = creats_for_insurance(
                await self.profit_insurance_list_by_code(sec_code_entity.securityCode),
                await self.cash_flow_insurance_list_by_code(sec_code_entity.securityCode),
                await self.roe_list_by_code(sec_code_entity.securityCode),
                sec_code_entity,
            )
        else:
            zqh = creats_for_bond(
                await self.profit_securities_list_by_code(sec_code_entity.securityCode),
                await self.cash_flow_securities_list_by_code(sec_code_entity.securityCode),
                await self.roe_list_by_code(sec_code_entity.securityCode),
                sec_code_entity,
            )
        await self._save_zqh_data(zqh, sec_code_entity)
        return zqh

    async def _save_zqh_data(self, zqh_list: List[ZqhFinEntity], sec_code_entity: SecCodeEntity):
        await self.fin_repo.save_many(
            [fin_to_entity(x, sec_code_entity.securityCode) for x in zqh_list]
        )

    async def _save_roe_mongo(self, yb_roe: YbRoeEntity, sec_code_entity: SecCodeEntity):
        if yb_roe is None:
            return
        kline_entity = await self.kline_service.kline_by_sec_code(sec_code_entity.securityCode)
        if not kline_entity or not kline_entity.klines:
            return
        kline_data = kline_entity.klines[0]
        price = Decimal(str(kline_data.close))
        roe_eps = YbRoeEpsEntity.model_validate(
            yb_roe.model_dump() if hasattr(yb_roe, "model_dump") else yb_roe.__dict__
        )
        roe_eps.code = sec_code_entity.securityCode
        roe_eps.name = sec_code_entity.securityNameAbbr
        roe_eps.id = f"{sec_code_entity.securityCode}{date_utils.today()}"

        if price >= roe_eps.pricehigher:
            roe_eps.value = "高估"
        elif price < roe_eps.pricehigher and price >= roe_eps.pricemiddle:
            roe_eps.value = "合理"
        elif price < roe_eps.pricemiddle and price >= roe_eps.pricelower:
            roe_eps.value = "偏低"
        else:
            roe_eps.value = "低估"

        if price >= roe_eps.pricequarterhigher:
            roe_eps.valuequarter = "高估"
        elif price < roe_eps.pricequarterhigher and price >= roe_eps.pricequartermiddle:
            roe_eps.valuequarter = "合理"
        elif price < roe_eps.pricequartermiddle and price >= roe_eps.pricequarterlower:
            roe_eps.valuequarter = "偏低"
        else:
            roe_eps.valuequarter = "低估"

        await self.roe_eps_repo.save(roe_eps)

    async def _get_yb_roe_entity(self, force: int, sec_code_entity: SecCodeEntity):
        listing_date = sec_code_entity.listingDate
        if (
            date_utils.period_to_next_date(
                date_utils.format_full_date_with_slip(listing_date), date_utils.today(), 3
            )
            < 4
        ):
            return None
        kline_entity = await self.kline_service.kline_by_sec_code(sec_code_entity.securityCode)
        if not kline_entity or not kline_entity.klines:
            return None
        profit_entities = []
        if sec_code_entity.orgTypeCode == UniversalTypeCode:
            profit_entities = [
                await self.profit_universal_list_by_code(sec_code_entity.securityCode)
            ]
        elif sec_code_entity.orgTypeCode == BankTypeCode:
            profit_entities = [await self.profit_bank_list_by_code(sec_code_entity.securityCode)]
        elif sec_code_entity.orgTypeCode == InsuranceTypeCode:
            profit_entities = [
                await self.profit_insurance_list_by_code(sec_code_entity.securityCode)
            ]
        elif sec_code_entity.orgTypeCode == SecuritiesTypeCode:
            profit_list = await self.profit_securities_list_by_code(sec_code_entity.securityCode)
            profit_entities = [profit_list]

        mapper = YbEpsEntityMapper()
        return mapper.creat(profit_entities, sec_code_entity, kline_entity.klines, force)

    async def _write_roe_kline_cells(self, ws, sec_code_entity: SecCodeEntity) -> None:
        try:
            kline_entity = await self.kline_service.kline_by_sec_code(sec_code_entity.securityCode)
        except Exception:
            return
        if not kline_entity or not kline_entity.klines:
            return
        latest = kline_entity.klines[0]
        date_value = self._normalize_date_value(getattr(latest, "date", None))
        close_value = getattr(latest, "close", None)
        if date_value is not None:
            ws["J13"].value = date_value
        if close_value is not None:
            ws["J14"].value = close_value

    async def _get_cash_flow_score_entities(self, dates: List, sec_code_entity: SecCodeEntity):
        items = []
        for d in dates:
            entity = await self.cash_universal_repo.find_by_id(
                f"{sec_code_entity.securityCode}{date_utils.date_to_string(d)}"
            )
            if entity:
                items.append(cash_score_to_entity(entity))
        return items

    async def get_score_entities(self, sec_code_entity: SecCodeEntity) -> List[ScoreEntity]:
        if sec_code_entity.orgTypeCode == UniversalTypeCode:
            return self._fin_score_universal(
                await self.profit_universal_list_by_code(sec_code_entity.securityCode),
                await self.cash_flow_universal_list_by_code(sec_code_entity.securityCode),
                await self.assets_universal_list_by_code(sec_code_entity.securityCode),
                await self.roe_list_by_code(sec_code_entity.securityCode),
                sec_code_entity,
            )
        if sec_code_entity.orgTypeCode == BankTypeCode:
            return self._fin_score_bank(
                await self.profit_bank_list_by_code(sec_code_entity.securityCode),
                await self.cash_flow_bank_list_by_code(sec_code_entity.securityCode),
                await self.assets_bank_list_by_code(sec_code_entity.securityCode),
                await self.roe_list_by_code(sec_code_entity.securityCode),
                sec_code_entity,
            )
        if sec_code_entity.orgTypeCode == SecuritiesTypeCode:
            return self._fin_score_securities(
                await self.profit_securities_list_by_code(sec_code_entity.securityCode),
                await self.cash_flow_securities_list_by_code(sec_code_entity.securityCode),
                await self.assets_securities_list_by_code(sec_code_entity.securityCode),
                await self.roe_list_by_code(sec_code_entity.securityCode),
                sec_code_entity,
            )
        return self._fin_score_insurance(
            await self.profit_insurance_list_by_code(sec_code_entity.securityCode),
            await self.cash_flow_insurance_list_by_code(sec_code_entity.securityCode),
            await self.assets_insurance_list_by_code(sec_code_entity.securityCode),
            await self.roe_list_by_code(sec_code_entity.securityCode),
            sec_code_entity,
        )

    def _fin_score_universal(
        self, profit_list, cash_list, assets_list, roe_list, sec_code_entity
    ) -> List[ScoreEntity]:
        profit = sorted(profit_list, key=lambda x: x.reportDate, reverse=True)
        cash = sorted(cash_list, key=lambda x: x.reportDate, reverse=True)
        assets = sorted(assets_list, key=lambda x: x.reportDate, reverse=True)
        roe = sorted(roe_list, key=lambda x: x.reportDate, reverse=True)
        scores = []
        for p in profit:
            cf = next((x for x in cash if x.id == p.id), None)
            a = next((x for x in assets if x.id == p.id), None)
            r = next((x for x in roe if x.id == p.id), None)
            if cf and a and r:
                scores.append(self._score_entity(p, a, cf, r, sec_code_entity))
        return scores

    def _fin_score_bank(
        self, profit_list, cash_list, assets_list, roe_list, sec_code_entity
    ) -> List[ScoreEntity]:
        profit = sorted(profit_list, key=lambda x: x.reportDate, reverse=True)
        cash = sorted(cash_list, key=lambda x: x.reportDate, reverse=True)
        assets = sorted(assets_list, key=lambda x: x.reportDate, reverse=True)
        roe = sorted(roe_list, key=lambda x: x.reportDate, reverse=True)
        scores = []
        for p in profit:
            cf = next((x for x in cash if x.id == p.id), None)
            a = next((x for x in assets if x.id == p.id), None)
            r = next((x for x in roe if x.id == p.id), None)
            if cf and a and r:
                scores.append(self._score_entity_bank(p, a, cf, r, sec_code_entity))
        return scores

    def _fin_score_insurance(
        self, profit_list, cash_list, assets_list, roe_list, sec_code_entity
    ) -> List[ScoreEntity]:
        profit = sorted(profit_list, key=lambda x: x.reportDate, reverse=True)
        cash = sorted(cash_list, key=lambda x: x.reportDate, reverse=True)
        assets = sorted(assets_list, key=lambda x: x.reportDate, reverse=True)
        roe = sorted(roe_list, key=lambda x: x.reportDate, reverse=True)
        scores = []
        for p in profit:
            cf = next((x for x in cash if x.id == p.id), None)
            a = next((x for x in assets if x.id == p.id), None)
            r = next((x for x in roe if x.id == p.id), None)
            if cf and a and r:
                scores.append(self._score_entity_bank(p, a, cf, r, sec_code_entity))
        return scores

    def _fin_score_securities(
        self, profit_list, cash_list, assets_list, roe_list, sec_code_entity
    ) -> List[ScoreEntity]:
        profit = sorted(profit_list, key=lambda x: x.reportDate, reverse=True)
        cash = sorted(cash_list, key=lambda x: x.reportDate, reverse=True)
        assets = sorted(assets_list, key=lambda x: x.reportDate, reverse=True)
        roe = sorted(roe_list, key=lambda x: x.reportDate, reverse=True)
        scores = []
        for p in profit:
            cf = next((x for x in cash if x.id == p.id), None)
            a = next((x for x in assets if x.id == p.id), None)
            r = next((x for x in roe if x.id == p.id), None)
            if cf and a and r:
                scores.append(self._score_entity_bank(p, a, cf, r, sec_code_entity))
        return scores

    def _score_entity(self, profit, assets, cash_flow, roe, sec_code_entity) -> ScoreEntity:
        s = ScoreEntity()
        s.id = f"{sec_code_entity.securityCode}{profit.reportDate.replace(' 00:00:00', '')}"
        s.code = sec_code_entity.securityCode
        s.name = sec_code_entity.securityNameAbbr
        s.date = profit.reportDate

        s.grossProfitPer = finance_utils.gross_profit(profit, sec_code_entity)
        s.operatProfitPer = finance_utils.operat_profit(profit, sec_code_entity)
        s.netProfitPer = finance_utils.net_profit(profit, sec_code_entity)
        s.netAssetsWeightPer = roe.roe

        s.liabilPer = finance_utils.liabil_per(assets)
        s.assetsPer = finance_utils.assets_per(assets)
        s.stockPer = finance_utils.stock_per(assets, profit, sec_code_entity)
        s.accountsReceivablePer = finance_utils.accounts_receivable_per(
            assets, profit, sec_code_entity
        )

        s.netCashFlowFromOperatingActivities = (
            finance_utils.net_cash_flow_from_operating_activities(cash_flow)
        )
        s.netCashFlowFromInvestmentActivities = (
            finance_utils.net_cash_flow_from_investment_activities(cash_flow)
        )
        s.netCashFlowFromFinancingActivities = (
            finance_utils.net_cash_flow_from_financing_activities(cash_flow)
        )

        score = (
            GrossProfitScore.includeStart(s.grossProfitPer)
            + OperateProfitScore.includeStart(s.operatProfitPer)
            + NetProfitScore.includeStart(s.netProfitPer)
            + NetAssetsWeightScore.includeStart(s.netAssetsWeightPer)
            + LiabilScore.includeStart(s.liabilPer)
            + AssetsScore.includeStart(s.assetsPer)
            + StockScore.includeStart(s.stockPer)
            + AccountsRecivableScore.includeStart(s.accountsReceivablePer)
            + CashFlowScore.netCashFlowFromOperatingActivities(s.netCashFlowFromOperatingActivities)
            + CashFlowScore.netCashFlowFromInvestmentActivities(
                s.netCashFlowFromInvestmentActivities
            )
            + CashFlowScore.netCashFlowFromFinancingActivities(s.netCashFlowFromFinancingActivities)
        )

        s.score = score
        s.result = ResultScore.includeStart(score)

        s.grossProfitScore = GrossProfitScore.includeStartString(s.grossProfitPer)
        s.operatProfitScore = OperateProfitScore.includeStartString(s.operatProfitPer)
        s.netProfitScore = NetProfitScore.includeStartString(s.netProfitPer)
        s.netAssetsWeightScore = NetAssetsWeightScore.includeStartString(s.netAssetsWeightPer)
        s.liabilScore = LiabilScore.includeStartString(s.liabilPer)
        s.assetsScore = AssetsScore.includeStartString(s.assetsPer)
        s.stockScore = StockScore.includeStartString(s.stockPer)
        s.accountsReceivableScore = AccountsRecivableScore.includeStartString(
            s.accountsReceivablePer
        )
        s.netCashFlowFromOperatingActivitiesScore = (
            CashFlowScore.netCashFlowFromOperatingActivitiesString(
                s.netCashFlowFromOperatingActivities
            )
        )
        s.netCashFlowFromInvestmentActivitiesScore = (
            CashFlowScore.netCashFlowFromInvestmentActivitiesString(
                s.netCashFlowFromInvestmentActivities
            )
        )
        s.netCashFlowFromFinancingActivitiesScore = (
            CashFlowScore.netCashFlowFromFinancingActivitiesString(
                s.netCashFlowFromFinancingActivities
            )
        )
        return s

    def _score_entity_bank(self, profit, assets, cash_flow, roe, sec_code_entity) -> ScoreEntity:
        s = ScoreEntity()
        s.id = f"{sec_code_entity.securityCode}{profit.reportDate.replace(' 00:00:00', '')}"
        s.code = sec_code_entity.securityCode
        s.name = sec_code_entity.securityNameAbbr
        s.date = profit.reportDate

        s.grossProfitPer = finance_utils.gross_profit(profit, sec_code_entity)
        s.operatProfitPer = finance_utils.operat_profit(profit, sec_code_entity)
        s.netProfitPer = finance_utils.net_profit(profit, sec_code_entity)
        s.netAssetsWeightPer = roe.roe

        s.netCashFlowFromOperatingActivities = (
            finance_utils.net_cash_flow_from_operating_activities(cash_flow)
        )
        s.netCashFlowFromInvestmentActivities = (
            finance_utils.net_cash_flow_from_investment_activities(cash_flow)
        )
        s.netCashFlowFromFinancingActivities = (
            finance_utils.net_cash_flow_from_financing_activities(cash_flow)
        )

        score = (
            GrossProfitBankScore.includeStart(s.grossProfitPer)
            + OperateProfitBankScore.includeStart(s.operatProfitPer)
            + NetProfitBankScore.includeStart(s.netProfitPer)
            + NetAssetsWeightBankScore.includeStart(s.netAssetsWeightPer)
            + CashFlowScore.netCashFlowFromOperatingActivities(s.netCashFlowFromOperatingActivities)
            + CashFlowScore.netCashFlowFromInvestmentActivities(
                s.netCashFlowFromInvestmentActivities
            )
            + CashFlowScore.netCashFlowFromFinancingActivities(s.netCashFlowFromFinancingActivities)
        )
        s.score = score
        s.result = ResultScore.includeStart(score)
        s.grossProfitScore = GrossProfitBankScore.includeStartString(s.grossProfitPer)
        s.operatProfitScore = OperateProfitBankScore.includeStartString(s.operatProfitPer)
        s.netProfitScore = NetProfitBankScore.includeStartString(s.netProfitPer)
        s.netAssetsWeightScore = NetAssetsWeightBankScore.includeStartString(s.netAssetsWeightPer)
        s.netCashFlowFromOperatingActivitiesScore = (
            CashFlowScore.netCashFlowFromOperatingActivitiesString(
                s.netCashFlowFromOperatingActivities
            )
        )
        s.netCashFlowFromInvestmentActivitiesScore = (
            CashFlowScore.netCashFlowFromInvestmentActivitiesString(
                s.netCashFlowFromInvestmentActivities
            )
        )
        s.netCashFlowFromFinancingActivitiesScore = (
            CashFlowScore.netCashFlowFromFinancingActivitiesString(
                s.netCashFlowFromFinancingActivities
            )
        )
        return s

    async def has_fin_data(self, sec_code_entity: SecCodeEntity, date: str) -> bool:
        key = f"{sec_code_entity.securityCode}{date}"
        if sec_code_entity.orgTypeCode == UniversalTypeCode:
            return (
                (await self.profit_universal_repo.find_by_id(key)) is not None
                and (await self.assets_universal_repo.find_by_id(key)) is not None
                and (await self.cash_universal_repo.find_by_id(key)) is not None
                and (await self.dupond_repo.find_by_id(key)) is not None
            )
        if sec_code_entity.orgTypeCode == BankTypeCode:
            return (
                (await self.profit_bank_repo.find_by_id(key)) is not None
                and (await self.assets_bank_repo.find_by_id(key)) is not None
                and (await self.cash_bank_repo.find_by_id(key)) is not None
                and (await self.dupond_repo.find_by_id(key)) is not None
            )
        if sec_code_entity.orgTypeCode == InsuranceTypeCode:
            return (
                (await self.profit_insurance_repo.find_by_id(key)) is not None
                and (await self.assets_insurance_repo.find_by_id(key)) is not None
                and (await self.cash_insurance_repo.find_by_id(key)) is not None
                and (await self.dupond_repo.find_by_id(key)) is not None
            )
        if sec_code_entity.orgTypeCode == SecuritiesTypeCode:
            return (
                (await self.profit_securities_repo.find_by_id(key)) is not None
                and (await self.assets_securities_repo.find_by_id(key)) is not None
                and (await self.cash_securities_repo.find_by_id(key)) is not None
                and (await self.dupond_repo.find_by_id(key)) is not None
            )
        return False

    async def has_fin_data_any(self, sec_code_entity: SecCodeEntity) -> bool:
        if sec_code_entity.orgTypeCode == UniversalTypeCode:
            return (
                bool(await self.profit_universal_list_by_code(sec_code_entity.securityCode))
                and bool(await self.assets_universal_list_by_code(sec_code_entity.securityCode))
                and bool(await self.cash_flow_universal_list_by_code(sec_code_entity.securityCode))
                and bool(await self.roe_list_by_code(sec_code_entity.securityCode))
            )
        if sec_code_entity.orgTypeCode == BankTypeCode:
            return (
                bool(await self.profit_bank_list_by_code(sec_code_entity.securityCode))
                and bool(await self.assets_bank_list_by_code(sec_code_entity.securityCode))
                and bool(await self.cash_flow_bank_list_by_code(sec_code_entity.securityCode))
                and bool(await self.roe_list_by_code(sec_code_entity.securityCode))
            )
        if sec_code_entity.orgTypeCode == InsuranceTypeCode:
            return (
                bool(await self.profit_insurance_list_by_code(sec_code_entity.securityCode))
                and bool(await self.assets_insurance_list_by_code(sec_code_entity.securityCode))
                and bool(await self.cash_flow_insurance_list_by_code(sec_code_entity.securityCode))
                and bool(await self.roe_list_by_code(sec_code_entity.securityCode))
            )
        if sec_code_entity.orgTypeCode == SecuritiesTypeCode:
            return (
                bool(await self.profit_securities_list_by_code(sec_code_entity.securityCode))
                and bool(await self.assets_securities_list_by_code(sec_code_entity.securityCode))
                and bool(await self.cash_flow_securities_list_by_code(sec_code_entity.securityCode))
                and bool(await self.roe_list_by_code(sec_code_entity.securityCode))
            )
        return False

    def _write_list(
        self,
        ws,
        items: List,
        start_row: int = 1,
        include_header: bool = True,
        normalize_dates: bool = False,
    ):
        if not items:
            return
        headers = (
            list(items[0].model_dump().keys())
            if hasattr(items[0], "model_dump")
            else list(items[0].__dict__.keys())
        )
        row = start_row
        if include_header:
            for c, h in enumerate(headers, start=1):
                ws.cell(row=row, column=c, value=h)
            row += 1
        for item in items:
            data = item.model_dump() if hasattr(item, "model_dump") else item.__dict__
            for c, h in enumerate(headers, start=1):
                value = data.get(h)
                if normalize_dates:
                    value = self._normalize_date_value(value)
                ws.cell(row=row, column=c, value=value)
            row += 1

    @staticmethod
    def _entity_to_dict(entity) -> dict:
        data = entity.model_dump() if hasattr(entity, "model_dump") else entity.__dict__
        extra = getattr(entity, "__pydantic_extra__", None)
        if isinstance(extra, dict) and extra:
            data = {**data, **extra}
        return data

    @staticmethod
    def _extract_year(report_date: str | None) -> str | None:
        if not report_date:
            return None
        import re
        m = re.match(r"(\d{4})", report_date)
        return m.group(1) if m else None

    async def _calc_dividend_yield_entities(self, sec_code_entity: SecCodeEntity) -> List[DividendYieldEntity]:
        sec_code = sec_code_entity.securityCode
        sec_name = sec_code_entity.securityNameAbbr

        sb_records = await self.share_bonus_list_by_code(sec_code)
        if not sb_records:
            return []

        year_sb: dict[str, list[dict]] = {}
        for b in sb_records:
            data = self._entity_to_dict(b)
            year = self._extract_year(data.get("reportDate", ""))
            if not year:
                continue
            year_sb.setdefault(year, []).append(data)

        profit_list = await self._get_profit_list_by_org_type(sec_code_entity)
        profit_by_year: dict[str, dict] = {}
        for p in profit_list:
            pdata = self._entity_to_dict(p)
            rd = pdata.get("reportDate", "")
            year = self._extract_year(rd)
            if year and ("12-31" in rd):
                profit_by_year[year] = pdata

        kline_entity = await self.kline_service.kline_by_sec_code(sec_code)
        klines_sorted = []
        if kline_entity and kline_entity.klines:
            from datetime import datetime as dt
            def _parse_kline_date(d: str) -> str:
                for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
                    try:
                        return dt.strptime(str(d)[:10], fmt).date().isoformat()
                    except ValueError:
                        continue
                return ""
            parsed = []
            for k in kline_entity.klines:
                kd = _parse_kline_date(k.date)
                if kd:
                    parsed.append((kd, k))
            parsed.sort(key=lambda x: x[0])
            klines_sorted = parsed

        result: List[DividendYieldEntity] = []
        for year in sorted(year_sb.keys(), reverse=True):
            records = year_sb[year]
            total_pretax = sum(float(b.get("pretaxBonusRmb", 0) or 0) for b in records)
            dividend_per_share = total_pretax / 10.0

            sb_total_shares = None
            for b in records:
                ts = b.get("totalShares")
                if ts is not None:
                    sb_total_shares = float(ts)
                    break
            if sb_total_shares is None:
                continue

            profit_data = profit_by_year.get(year)
            if not profit_data:
                continue

            netprofit_val = profit_data.get("parentNetprofit") or profit_data.get("netprofit")
            if not netprofit_val:
                continue

            total_shares_raw = sb_total_shares
            netprofit_raw = float(netprofit_val)

            records_implemented = [b for b in records if b.get("assignProgress") == "实施分配"]
            records_planned = [b for b in records if b.get("assignProgress") != "实施分配"]

            ex_date = None
            for b in records_implemented:
                d = b.get("exDividendDate") or b.get("noticeDate")
                if d:
                    ex_date = str(d)[:10]
                    break

            ex_price = None
            if records_implemented and not records_planned:
                if ex_date and klines_sorted:
                    from datetime import datetime as dt2
                    ex_dt = None
                    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
                        try:
                            ex_dt = dt2.strptime(ex_date[:10], fmt).date()
                            break
                        except ValueError:
                            continue
                    if ex_dt:
                        ex_date_iso = ex_dt.isoformat()
                        for kd_str, k in klines_sorted:
                            if kd_str >= ex_date_iso:
                                ex_price = float(k.close or 0)
                                break

            if ex_price is None and klines_sorted:
                ex_price = float(klines_sorted[-1][1].close or 0)

            total_shares_yi = total_shares_raw / 1e8
            netprofit_yi = netprofit_raw / 1e8

            payout_ratio = None
            if netprofit_raw != 0:
                payout_ratio = (dividend_per_share * total_shares_raw) / netprofit_raw * 100

            yield_pct = None
            if ex_price and ex_price != 0:
                yield_pct = dividend_per_share / ex_price * 100

            entity = DividendYieldEntity(
                secCode=sec_code,
                secName=sec_name,
                reportYear=year,
                pretaxBonusRmb=total_pretax,
                dividendPerShare=round(dividend_per_share, 4),
                totalShares=round(total_shares_yi, 2),
                netprofit=round(netprofit_yi, 2),
                exDividendDate=ex_date,
                exDividendPrice=round(ex_price, 2) if ex_price else None,
                dividendYield=round(yield_pct, 2) if yield_pct else None,
                dividendPayoutRatio=round(payout_ratio, 2) if payout_ratio else None,
            )
            result.append(entity)

        years_with_dividends = sorted(year_sb.keys(), key=int)
        consecutive = 0
        if years_with_dividends:
            consecutive = 1
            for i in range(len(years_with_dividends) - 1):
                if int(years_with_dividends[i + 1]) - int(years_with_dividends[i]) == 1:
                    consecutive += 1
                else:
                    break

        for entity in result:
            entity.consecutiveYears = consecutive

        return result

    def _write_dividend_sheet(self, ws, items: List[DividendYieldEntity]):
        headers = [
            "分红年度", "每10股分红(元)", "每股分红(元)",
            "总股本(亿股)", "净利润(亿元)", "除息日/公告日", "除息日股价(元)",
            "股息率(%)", "分红率(%)", "连续分红年数",
        ]
        header_fill = PatternFill(fill_type="solid", fgColor="BFBFBF")
        header_align = Alignment(horizontal="center", vertical="center")
        header_font = Font(name="Microsoft YaHei", size=10, color="00008B", bold=True)
        ws.row_dimensions[1].height = 20

        col_yield = 8
        yield_header_font = Font(name="Microsoft YaHei", size=10, color="008000", bold=True)

        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = yield_header_font if c == col_yield else header_font
            cell.fill = header_fill
            cell.alignment = header_align

        if items:
            data_font = Font(name="Calibri", size=11, bold=False)
            data_align_right = Alignment(horizontal="right", vertical="center")
            data_align_center = Alignment(horizontal="center", vertical="center")
            for row_idx, item in enumerate(items, start=2):
                data = item.model_dump() if hasattr(item, "model_dump") else item.__dict__
                vals = [data.get(h) for h in [
                    "reportYear", "pretaxBonusRmb",
                    "dividendPerShare", "totalShares", "netprofit",
                    "exDividendDate", "exDividendPrice",
                    "dividendYield", "dividendPayoutRatio", "consecutiveYears",
                ]]
                for c, v in enumerate(vals, start=1):
                    cell = ws.cell(row=row_idx, column=c, value=v)
                    cell.font = data_font
                    cell.alignment = data_align_center if c in {1, 6} else data_align_right

            yield_font_green = Font(name="Calibri", size=11, bold=False, color="008000")
            yield_font_red = Font(name="Calibri", size=11, bold=False, color="FF0000")
            for row_idx, item in enumerate(items, start=2):
                yield_val = item.dividendYield
                cell = ws.cell(row=row_idx, column=col_yield)
                if yield_val is not None and yield_val > 4:
                    cell.font = yield_font_red
                else:
                    cell.font = yield_font_green

        self._apply_border_to_range(ws, min_row=1, max_row=1, min_col=1, max_col=len(headers))
        if items:
            self._auto_fit_columns(ws, min_row=1, max_row=1 + len(items), min_col=1, max_col=len(headers), padding=2)
        else:
            self._auto_fit_columns(ws, min_row=1, max_row=1, min_col=1, max_col=len(headers), padding=2)

    def _write_zqh_sheet(self, ws, items: List):
        # Write a two-row Chinese header to match the original Java EasyExcel output
        headers_row1 = [
            "日期/科目",
            "预测公司成长性指标",
            "预测公司成长性指标",
            "预测公司成长性指标",
            "预测公司成长性指标",
            "分析公司获利性指标",
            "分析公司获利性指标",
            "分析公司获利性指标",
            "分析公司获利性指标",
            "检视公司安全性指标",
            "检视公司安全性指标",
        ]
        headers_row2 = [
            "日期/科目",
            "营业收入(亿元)",
            "营收增长率%",
            "净利润(亿元)",
            "净利润增长率%",
            "营业毛利率%",
            "净利率%",
            "营业利润率%",
            "净资产收益率%",
            "经营现金流净额(亿元)",
            "长短期负债比",
        ]

        for c, v in enumerate(headers_row1, start=1):
            ws.cell(row=1, column=c, value=v)
        for c, v in enumerate(headers_row2, start=1):
            ws.cell(row=2, column=c, value=v)

        # Column widths based on Java @ColumnWidth annotations (default 10 if not specified)
        col_widths = [10, 20, 18, 20, 20, 20, 15, 20, 20, 20, 15]
        for i, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # Header row height based on @HeadRowHeight(20)
        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 20

        # Header font size based on @HeadFontStyle(fontHeightInPoints = 10)
        # Colors use Excel indexed palette (as in Java color indices)
        color_grow = Color(indexed=21)
        color_profit = Color(indexed=10)
        color_safe = Color(indexed=24)
        font_grow = Font(name="Microsoft YaHei", size=10, color=color_grow, bold=True)
        font_profit = Font(name="Microsoft YaHei", size=10, color=color_profit, bold=True)
        font_safe = Font(name="Microsoft YaHei", size=10, color=color_safe, bold=True)
        header_fill = PatternFill(fill_type="solid", fgColor="D9D9D9")
        header_align = Alignment(horizontal="center", vertical="center")

        # Apply header styles per group
        for col in range(1, 12):
            ws.cell(row=1, column=col).fill = header_fill
            ws.cell(row=2, column=col).fill = header_fill
            ws.cell(row=1, column=col).alignment = header_align
            ws.cell(row=2, column=col).alignment = header_align
        for col in range(2, 6):
            ws.cell(row=1, column=col).font = font_grow
            ws.cell(row=2, column=col).font = font_grow
        for col in range(6, 10):
            ws.cell(row=1, column=col).font = font_profit
            ws.cell(row=2, column=col).font = font_profit
        for col in range(10, 12):
            ws.cell(row=1, column=col).font = font_safe
            ws.cell(row=2, column=col).font = font_safe
        # Column 1 header uses default font size 10
        ws.cell(row=1, column=1).font = Font(name="Microsoft YaHei", size=10, bold=True)
        ws.cell(row=2, column=1).font = Font(name="Microsoft YaHei", size=10, bold=True)

        # Merge group headers to match the screenshot layout
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=5)
        ws.merge_cells(start_row=1, start_column=6, end_row=1, end_column=9)
        ws.merge_cells(start_row=1, start_column=10, end_row=1, end_column=11)

        # Write data starting from row 3
        self._write_list(ws, items, start_row=3, include_header=False, normalize_dates=True)

        # Apply content font colors per column (match @ContentFontStyle)
        if items:
            last_row = 2 + len(items)
            date_align = Alignment(horizontal="center", vertical="center")
            for row in range(3, last_row + 1):
                ws.cell(row=row, column=1).font = Font(
                    name="Calibri", size=11, color="404040", bold=False
                )
                ws.cell(row=row, column=1).alignment = date_align
                for col in range(2, 6):
                    ws.cell(row=row, column=col).font = Font(
                        name="Calibri", size=11, color=color_grow, bold=False
                    )
                for col in range(6, 10):
                    ws.cell(row=row, column=col).font = Font(
                        name="Calibri", size=11, color=color_profit, bold=False
                    )
                for col in range(10, 12):
                    ws.cell(row=row, column=col).font = Font(
                        name="Calibri", size=11, color=color_safe, bold=False
                    )

        # Add black borders only to header/title area
        self._apply_border_to_range(ws, min_row=1, max_row=2, min_col=1, max_col=11)
        if items:
            self._auto_fit_columns(
                ws, min_row=1, max_row=2 + len(items), min_col=1, max_col=11, padding=2
            )
        else:
            self._auto_fit_columns(ws, min_row=1, max_row=2, min_col=1, max_col=11, padding=2)

    def _write_cash_flow_sheet(self, ws, items: List):
        # Header row based on original Java ExcelProperty titles
        headers = [
            "日期",
            "代码",
            "名称",
            "经营现金流(亿元)",
            "投资现金流(亿元)",
            "筹资现金流(亿元)",
            "类型",
            "特征",
        ]
        header_fill = PatternFill(fill_type="solid", fgColor="BFBFBF")
        header_align = Alignment(horizontal="center", vertical="center")
        header_font = Font(name="Microsoft YaHei", size=10, color="00008B", bold=True)
        profit_font = Font(name="Microsoft YaHei", size=10, color="1F77B4", bold=True)
        structure_font = Font(name="Microsoft YaHei", size=10, color="D62728", bold=True)
        cash_font = Font(name="Microsoft YaHei", size=10, color="2CA02C", bold=True)
        profit_font = Font(name="Microsoft YaHei", size=10, color="1F77B4", bold=True)
        structure_font = Font(name="Microsoft YaHei", size=10, color="D62728", bold=True)
        cash_font = Font(name="Microsoft YaHei", size=10, color="2CA02C", bold=True)
        profit_font = Font(name="Microsoft YaHei", size=10, color="1F77B4", bold=True)
        structure_font = Font(name="Microsoft YaHei", size=10, color="D62728", bold=True)
        cash_font = Font(name="Microsoft YaHei", size=10, color="2CA02C", bold=True)
        profit_font = Font(name="Microsoft YaHei", size=10, color="1F77B4", bold=True)
        structure_font = Font(name="Microsoft YaHei", size=10, color="D62728", bold=True)
        cash_font = Font(name="Microsoft YaHei", size=10, color="2CA02C", bold=True)
        header_font_deep_green = Font(name="Microsoft YaHei", size=10, color="006400", bold=True)
        ws.row_dimensions[1].height = 20
        for c, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c, value=h)
            ws.cell(row=1, column=c).font = header_font_deep_green if c in {17, 19} else header_font
            ws.cell(row=1, column=c).fill = header_fill
            ws.cell(row=1, column=c).alignment = header_align
            # Column width based on title length + 2
            ws.column_dimensions[get_column_letter(c)].width = len(str(h)) + 2

        # Write data starting from row 2
        self._write_list(ws, items, start_row=2, include_header=False, normalize_dates=True)
        if items:
            data_font = Font(name="Calibri", size=11, bold=False)
            data_align = Alignment(horizontal="right", vertical="center")
            date_align = Alignment(horizontal="center", vertical="center")
            type_highlight_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
            type_highlight_align = Alignment(horizontal="center", vertical="center")
            type_highlight_values_red = {"妖精型", "老母鸡性", "蛮牛型", "奶牛型"}
            type_highlight_values_green = {"骗吃骗喝型", "混吃等死型", "赌徒型", "大出血型"}
            type_highlight_fill_red = PatternFill(fill_type="solid", fgColor="FF0000")
            type_highlight_fill_green = PatternFill(fill_type="solid", fgColor="00FF00")
            for row in ws.iter_rows(min_row=2, max_row=1 + len(items)):
                for cell in row:
                    if cell.value is not None:
                        if cell.column == 7:
                            text = str(cell.value)
                            if text in type_highlight_values_red:
                                cell.font = type_highlight_font
                                cell.fill = type_highlight_fill_red
                                cell.alignment = type_highlight_align
                            elif text in type_highlight_values_green:
                                cell.font = type_highlight_font
                                cell.fill = type_highlight_fill_green
                                cell.alignment = type_highlight_align
                            else:
                                cell.font = data_font
                                cell.alignment = date_align if cell.column == 1 else data_align
                        else:
                            cell.font = data_font
                            cell.alignment = date_align if cell.column == 1 else data_align

        # Add black borders only to header/title row
        self._apply_border_to_range(ws, min_row=1, max_row=1, min_col=1, max_col=len(headers))
        if items:
            self._auto_fit_columns(
                ws, min_row=1, max_row=1 + len(items), min_col=1, max_col=len(headers), padding=2
            )
        else:
            self._auto_fit_columns(
                ws, min_row=1, max_row=1, min_col=1, max_col=len(headers), padding=2
            )

    async def _build_fcf_entities(self, sec_code_entity: SecCodeEntity) -> List[FreeCashFlowEntity]:
        cash_items = await self.cash_flow_universal_list_by_code(sec_code_entity.securityCode)
        if not cash_items:
            return []

        profit_items = await self.profit_universal_list_by_code(sec_code_entity.securityCode)
        profit_by_date = {}
        for p in profit_items:
            d = str(p.id).replace(p.securityCode, "") if p.id else None
            if d:
                profit_by_date[d] = p

        result: List[FreeCashFlowEntity] = []
        for cf in cash_items:
            date = str(cf.id).replace(cf.securityCode, "") if cf.id else None
            if not date:
                continue

            e = FreeCashFlowEntity()
            e.date = date
            e.securityCode = cf.securityCode
            e.securityNameAbbr = cf.securityNameAbbr
            e.operatingCashFlow = num_utils.round_double(
                num_utils.string_to_double(getattr(cf, "netcashOperate", None)) / 10000 / 10000
            )
            e.capex = num_utils.round_double(
                num_utils.string_to_double(getattr(cf, "constructLongAsset", None)) / 10000 / 10000
            )
            e.freeCashFlow = finance_utils.free_cash_flow(cf)

            profit_match = profit_by_date.get(date)
            if profit_match:
                np_val = num_utils.string_to_double(getattr(profit_match, "parentNetprofit", None))
                np_yi = num_utils.round_double(np_val / 10000 / 10000)
                if np_yi and np_yi > 0:
                    e.fcfToNetProfit = num_utils.round_double(e.freeCashFlow / np_yi * 100)
                else:
                    e.fcfToNetProfit = 0.0
            else:
                e.fcfToNetProfit = 0.0

            result.append(e)
        return result

    def _write_fcf_sheet(self, ws, items: List[FreeCashFlowEntity]):
        headers = [
            "日期",
            "代码",
            "名称",
            "经营现金流(亿元)",
            "资本开支(亿元)",
            "自由现金流(亿元)",
            "FCF/净利润(%)",
        ]
        header_fill = PatternFill(fill_type="solid", fgColor="BFBFBF")
        header_align = Alignment(horizontal="center", vertical="center")
        header_font = Font(name="Microsoft YaHei", size=10, color="00008B", bold=True)
        ws.row_dimensions[1].height = 20
        for c, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c, value=h)
            ws.cell(row=1, column=c).font = header_font
            ws.cell(row=1, column=c).fill = header_fill
            ws.cell(row=1, column=c).alignment = header_align
            ws.column_dimensions[get_column_letter(c)].width = len(str(h)) + 2

        order = [
            "date",
            "securityCode",
            "securityNameAbbr",
            "operatingCashFlow",
            "capex",
            "freeCashFlow",
            "fcfToNetProfit",
        ]
        if not items:
            return
        row = 2
        data_font = Font(name="Calibri", size=11, bold=False)
        data_align = Alignment(horizontal="right", vertical="center")
        date_align = Alignment(horizontal="center", vertical="center")
        for item in items:
            data = item.model_dump() if hasattr(item, "model_dump") else item.__dict__
            for c, field in enumerate(order, start=1):
                value = data.get(field)
                ws.cell(row=row, column=c, value=value)
                cell = ws.cell(row=row, column=c)
                cell.font = data_font
                cell.alignment = date_align if c == 1 else data_align
            row += 1

        self._apply_border_to_range(ws, min_row=1, max_row=1, min_col=1, max_col=len(headers))
        if items:
            self._auto_fit_columns(
                ws, min_row=1, max_row=1 + len(items), min_col=1, max_col=len(headers), padding=2
            )
        else:
            self._auto_fit_columns(
                ws, min_row=1, max_row=1, min_col=1, max_col=len(headers), padding=2
            )

    def _write_score_sheet(self, ws, items: List, org_type: str | None = None):
        # Two-row header: keep existing titles as row 2, add grouped titles in row 1
        is_finance_related = org_type in {BankTypeCode, InsuranceTypeCode, SecuritiesTypeCode}
        if is_finance_related:
            headers_row2 = [
                "日期",
                "总分",
                "总分评级",
                "毛利率",
                "毛利率得分",
                "营业利润率",
                "营业利润率得分",
                "净利率",
                "净利率得分",
                "净资产收益率",
                "净资产收益率得分",
                "经营活动现金流量净额",
                "经营活动现金流量净额得分",
                "投资活动现金流量净额",
                "投资活动现金流量净额得分",
                "筹资活动现金流量净额",
                "筹资活动现金流量净额得分",
            ]
        else:
            headers_row2 = [
                "日期",
                "总分",
                "总分评级",
                "毛利率",
                "毛利率得分",
                "营业利润率",
                "营业利润率得分",
                "净利率",
                "净利率得分",
                "净资产收益率",
                "净资产收益率得分",
                "资产负债比率",
                "资产负债比率得分",
                "流动资产/总资产",
                "流动资产/总资产得分",
                "存货周转率",
                "存货周转率得分",
                "应收账款周转率",
                "应收账款周转率得分",
                "经营活动现金流量净额",
                "经营活动现金流量净额得分",
                "投资活动现金流量净额",
                "投资活动现金流量净额得分",
                "筹资活动现金流量净额",
                "筹资活动现金流量净额得分",
            ]
        header_fill = PatternFill(fill_type="solid", fgColor="BFBFBF")
        header_align = Alignment(horizontal="center", vertical="center")
        header_font = Font(name="Microsoft YaHei", size=10, color="00008B", bold=True)
        profit_font = Font(name="Microsoft YaHei", size=10, color="1F77B4", bold=True)
        structure_font = Font(name="Microsoft YaHei", size=10, color="D62728", bold=True)
        cash_font = Font(name="Microsoft YaHei", size=10, color="2CA02C", bold=True)
        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 20

        # Row 1 group headers
        ws.cell(row=1, column=1, value="日期")
        ws.cell(row=1, column=2, value="综合评级")
        ws.cell(row=1, column=4, value="获利能力")
        if not is_finance_related:
            ws.cell(row=1, column=12, value="财务结构")
            ws.cell(row=1, column=20, value="现金流量")
        else:
            ws.cell(row=1, column=12, value="现金流量")
        for c in range(1, len(headers_row2) + 1):
            ws.cell(row=1, column=c).font = header_font
            ws.cell(row=1, column=c).fill = header_fill
            ws.cell(row=1, column=c).alignment = header_align

        # Row 2 titles (existing headers)
        for c, h in enumerate(headers_row2, start=1):
            if c == 1:
                continue  # merged "日期" occupies rows 1-2
            ws.cell(row=2, column=c, value=h)
            ws.cell(row=2, column=c).font = header_font
            ws.cell(row=2, column=c).fill = header_fill
            ws.cell(row=2, column=c).alignment = header_align
            # Column width based on title length + 2
            ws.column_dimensions[get_column_letter(c)].width = len(str(h)) + 2
        # Ensure date column has a reasonable width even when no data
        ws.column_dimensions[get_column_letter(1)].width = len("日期") + 2

        # Merge group headers across their columns
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
        ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=11)
        if not is_finance_related:
            ws.merge_cells(start_row=1, start_column=12, end_row=1, end_column=19)
            ws.merge_cells(start_row=1, start_column=20, end_row=1, end_column=25)
        else:
            ws.merge_cells(start_row=1, start_column=12, end_row=1, end_column=17)

        # Apply group header colors (row 1 + row 2)
        profit_cols = set(range(4, 12))
        if is_finance_related:
            structure_cols = set()
            cash_cols = set(range(12, 18))
        else:
            structure_cols = set(range(12, 20))
            cash_cols = set(range(20, 26))
        for col in profit_cols:
            ws.cell(row=1, column=col).font = profit_font
            ws.cell(row=2, column=col).font = profit_font
        for col in structure_cols:
            ws.cell(row=1, column=col).font = structure_font
            ws.cell(row=2, column=col).font = structure_font
        for col in cash_cols:
            ws.cell(row=1, column=col).font = cash_font
            ws.cell(row=2, column=col).font = cash_font

        # Ensure data columns match the header order above
        if is_finance_related:
            order = [
                "date",
                "score",
                "result",
                "grossProfitPer",
                "grossProfitScore",
                "operatProfitPer",
                "operatProfitScore",
                "netProfitPer",
                "netProfitScore",
                "netAssetsWeightPer",
                "netAssetsWeightScore",
                "netCashFlowFromOperatingActivities",
                "netCashFlowFromOperatingActivitiesScore",
                "netCashFlowFromInvestmentActivities",
                "netCashFlowFromInvestmentActivitiesScore",
                "netCashFlowFromFinancingActivities",
                "netCashFlowFromFinancingActivitiesScore",
            ]
        else:
            order = [
                "date",
                "score",
                "result",
                "grossProfitPer",
                "grossProfitScore",
                "operatProfitPer",
                "operatProfitScore",
                "netProfitPer",
                "netProfitScore",
                "netAssetsWeightPer",
                "netAssetsWeightScore",
                "liabilPer",
                "liabilScore",
                "assetsPer",
                "assetsScore",
                "stockPer",
                "stockScore",
                "accountsReceivablePer",
                "accountsReceivableScore",
                "netCashFlowFromOperatingActivities",
                "netCashFlowFromOperatingActivitiesScore",
                "netCashFlowFromInvestmentActivities",
                "netCashFlowFromInvestmentActivitiesScore",
                "netCashFlowFromFinancingActivities",
                "netCashFlowFromFinancingActivitiesScore",
            ]
        if not items:
            return
        row = 3
        data_font = Font(name="Calibri", size=11, bold=False)
        data_profit_font = Font(name="Calibri", size=11, bold=False, color="1F77B4")
        data_structure_font = Font(name="Calibri", size=11, bold=False, color="D62728")
        data_cash_font = Font(name="Calibri", size=11, bold=False, color="2CA02C")
        data_align = Alignment(horizontal="right", vertical="center")
        date_align = Alignment(horizontal="center", vertical="center")
        score_red_font = Font(name="Calibri", size=11, bold=False, color="FF0000")
        score_sky_font = Font(name="Calibri", size=11, bold=False, color="87CEFA")
        score_orange_font = Font(name="Calibri", size=11, bold=False, color="FFA500")
        score_grass_font = Font(name="Calibri", size=11, bold=False, color="7CFC00")
        score_deep_green_font = Font(name="Calibri", size=11, bold=False, color="006400")
        score_deep_blue_font = Font(name="Calibri", size=11, bold=False, color="00008B")
        rating_good_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        rating_mid_font = Font(name="Microsoft YaHei", size=11, bold=True, color="000000")
        rating_watch_font = Font(name="Microsoft YaHei", size=11, bold=True, color="006400")
        rating_good_fill = PatternFill(fill_type="solid", fgColor="FF0000")
        rating_mid_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
        rating_watch_fill = PatternFill(fill_type="solid", fgColor="00FF00")
        red_cols = {3}  # 总分评级
        sky_cols = {5, 7, 9, 11}  # 各类得分
        if is_finance_related:
            orange_cols = set()
            grass_cols = set()
            deep_blue_cols = {13, 15, 17}  # 三个现金流量得分
        else:
            orange_cols = {13, 15}  # 资产负债比率得分, 流动资产/总资产得分
            grass_cols = {17, 19}  # 存货周转率得分, 应收账款周转率得分
            deep_blue_cols = {21, 23, 25}  # 三个现金流量得分
        if is_finance_related:
            profit_cols_data = set(range(4, 12))
            structure_cols_data = set()
            cash_cols_data = set(range(12, 18))
        else:
            profit_cols_data = set(range(4, 12))
            structure_cols_data = set(range(12, 20))
            cash_cols_data = set(range(20, 26))
        for item in items:
            data = item.model_dump() if hasattr(item, "model_dump") else item.__dict__
            for c, key in enumerate(order, start=1):
                value = (
                    self._normalize_date_value(data.get(key)) if key == "date" else data.get(key)
                )
                ws.cell(row=row, column=c, value=value)
                if value is not None:
                    if c == 3:
                        rating_text = str(value)
                        if rating_text == "优等":
                            ws.cell(row=row, column=c).font = rating_good_font
                            ws.cell(row=row, column=c).fill = rating_good_fill
                        elif rating_text == "中等":
                            ws.cell(row=row, column=c).font = rating_mid_font
                            ws.cell(row=row, column=c).fill = rating_mid_fill
                        elif rating_text == "观望":
                            ws.cell(row=row, column=c).font = rating_watch_font
                            ws.cell(row=row, column=c).fill = rating_watch_fill
                        else:
                            ws.cell(row=row, column=c).font = data_font
                    elif c in profit_cols_data:
                        ws.cell(row=row, column=c).font = data_profit_font
                    elif c in structure_cols_data:
                        ws.cell(row=row, column=c).font = data_structure_font
                    elif c in cash_cols_data:
                        ws.cell(row=row, column=c).font = data_cash_font
                    elif c in sky_cols:
                        ws.cell(row=row, column=c).font = score_sky_font
                    elif c in orange_cols:
                        ws.cell(row=row, column=c).font = score_orange_font
                    elif c in grass_cols:
                        ws.cell(row=row, column=c).font = score_deep_green_font
                    elif c in deep_blue_cols:
                        ws.cell(row=row, column=c).font = score_deep_blue_font
                    else:
                        ws.cell(row=row, column=c).font = data_font
                    if c == 1 or c == 3:
                        ws.cell(row=row, column=c).alignment = date_align
                    else:
                        ws.cell(row=row, column=c).alignment = data_align
            row += 1

        # Add black borders only to header/title rows
        self._apply_border_to_range(ws, min_row=1, max_row=2, min_col=1, max_col=len(headers_row2))
        if items:
            self._auto_fit_columns(
                ws,
                min_row=1,
                max_row=2 + len(items),
                min_col=1,
                max_col=len(headers_row2),
                padding=2,
            )
        else:
            self._auto_fit_columns(
                ws, min_row=1, max_row=2, min_col=1, max_col=len(headers_row2), padding=2
            )

    def _apply_border_to_range(self, ws, min_row: int, max_row: int, min_col: int, max_col: int):
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if cell.value is not None:
                    cell.border = border

    def _normalize_date_value(self, value):
        if value is None:
            return value
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, str) and value.endswith(" 00:00:00"):
            return value.replace(" 00:00:00", "")
        return value

    def _auto_fit_columns(
        self, ws, min_row: int, max_row: int, min_col: int, max_col: int, padding: int = 2
    ):
        for col in range(min_col, max_col + 1):
            max_len = 0
            for row in range(min_row, max_row + 1):
                v = ws.cell(row=row, column=col).value
                if v is None:
                    continue
                text = str(v)
                display_len = 0
                for ch in text:
                    display_len += 2 if ord(ch) > 127 else 1
                if display_len > max_len:
                    max_len = display_len
            ws.column_dimensions[get_column_letter(col)].width = max_len + padding

    def _write_dict(self, ws, data: dict):
        def _is_writable_cell(r: int, c: int) -> bool:
            cell = ws.cell(row=r, column=c)
            if cell.coordinate not in ws.merged_cells:
                return True
            for mr in ws.merged_cells.ranges:
                if cell.coordinate in mr:
                    return mr.min_row == r and mr.min_col == c
            return True

        row = 1
        for k, v in data.items():
            while not (_is_writable_cell(row, 1) and _is_writable_cell(row, 2)):
                row += 1
            ws.cell(row=row, column=1, value=k)
            ws.cell(row=row, column=2, value=str(v))
            row += 1

    def _fill_template(self, ws, data: dict):
        if not data:
            return
        safe = {str(k): "" if v is None else str(v) for k, v in data.items()}
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str) or "{" not in v or "}" not in v:
                    continue
                new_v = v
                for k, val in safe.items():
                    token = "{" + k + "}"
                    if token in new_v:
                        new_v = new_v.replace(token, val)
                if new_v != v:
                    cell.value = new_v
