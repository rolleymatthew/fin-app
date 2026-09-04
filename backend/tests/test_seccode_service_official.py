import os

import openpyxl
import pytest

from app.clients.exchange_base import OfficialStock
from app.clients.sse_equity import SseEquityClient
from app.clients.szse_xlsx import SzseXlsxClient


# --------------------------------------------------------------------- #
# SSE equity row mapper
# --------------------------------------------------------------------- #
def test_sse_equity_row_to_official_e110():
    out = SseEquityClient._row_to_official(["600000", "浦发银行", "E110    "])
    assert out == OfficialStock(
        code="600000", name="浦发银行", market="SH",
        listingDate=None, listingState="0", type=None,
        raw={"code": "600000", "name": "浦发银行", "tradephase": "E110    "},
    )


def test_sse_equity_row_to_official_p010_marks_suspension():
    out = SseEquityClient._row_to_official(["603221", "爱丽家居", "P010    "])
    assert out is not None
    assert out.code == "603221"
    assert out.name == "爱丽家居"
    assert out.listingState == "2"
    assert out.market == "SH"


def test_sse_equity_row_to_official_invalid_code_returns_none():
    assert SseEquityClient._row_to_official(["abc", "x", "E110    "]) is None
    assert SseEquityClient._row_to_official(["", "x", "E110    "]) is None
    assert SseEquityClient._row_to_official(["600000", "", "E110    "]) is None


def test_sse_equity_row_to_official_unknown_tradephase_marks_suspension():
    """非 E110 视为非常规（含 P010/H000/...），保守标 '2'。"""
    out = SseEquityClient._row_to_official(["600999", "某停牌", "H000    "])
    assert out is not None
    assert out.listingState == "2"


# --------------------------------------------------------------------- #
# SZSE xlsx row mapper
# --------------------------------------------------------------------- #
SZSE_HEADER_MAP = {
    0: "板块",
    4: "A股代码",
    5: "A股简称",
    6: "A股上市日期",
}


def test_szse_xlsx_row_to_official_basic():
    row = [
        "主板", "平安银行股份有限公司", "Ping An Bank Co., Ltd.",
        "广东省深圳市", "000001", "平安银行", "1991-04-03",
        "19405918198", "19405684991", "", "", "", "0", "0",
        "华南", "广东", "深圳市", "J 金融业", "bank.pingan.com", "-", "-", "-",
    ]
    out = SzseXlsxClient._row_to_official(row, SZSE_HEADER_MAP)
    assert out == OfficialStock(
        code="000001", name="平安银行", market="SZ",
        listingDate="1991-04-03", listingState="0", type="主板",
        raw={k: row[i] for i, k in SZSE_HEADER_MAP.items()},
    )


def test_szse_xlsx_row_to_official_chuangyeban():
    row = [
        "创业板", "宁德时代", "CATL", "福建宁德", "300750", "宁德时代", "2018-06-11",
        "4400000000", "4400000000", "", "", "", "0", "0",
        "华东", "福建", "宁德市", "C 制造业", "www.catl.com", "-", "-", "-",
    ]
    out = SzseXlsxClient._row_to_official(row, SZSE_HEADER_MAP)
    assert out is not None
    assert out.market == "SZ"
    assert out.type == "创业板"
    assert out.listingDate == "2018-06-11"


def test_szse_xlsx_row_to_official_invalid_code_returns_none():
    row = [
        "主板", "x", "x", "x", "abc", "x", "x", "", "", "", "", "",
        "", "", "", "", "", "", "", "", "", "", "", "",
    ]
    assert SzseXlsxClient._row_to_official(row, SZSE_HEADER_MAP) is None


def test_szse_xlsx_parse_fixture():
    """解析真实格式的 fixture，5 行数据。"""
    fixture = os.path.join(os.path.dirname(__file__), "fixtures", "szse_a_stock.xlsx")
    wb = openpyxl.load_workbook(fixture, data_only=True)
    ws = wb.active
    header = [c.value for c in ws[1]]
    header_map = {i: h for i, h in enumerate(header)}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        stock = SzseXlsxClient._row_to_official(list(row), header_map)
        if stock is not None:
            rows.append(stock)
    codes = [r.code for r in rows]
    assert codes == ["000001", "000002", "000006", "300750", "000010"]
    assert all(r.market == "SZ" for r in rows)
    assert all(r.listingState == "0" for r in rows)
    assert rows[3].type == "创业板"
    assert rows[0].listingDate == "1991-04-03"


_SZSE_HEADER_MAP = {0: "板块", 4: "A股代码", 5: "A股简称", 6: "A股上市日期"}


def _szse_row(code: str, name: str, date: str = "1991-04-03") -> list:
    return [
        "主板", "x", "x", "x", code, name, date,
        "", "", "", "", "", "", "", "", "", "", "", "", "", "", "",
    ]


@pytest.mark.asyncio
async def test_fetch_official_uses_two_sources(monkeypatch):
    from app.services.seccode_service import SecCodeService
    svc = SecCodeService.__new__(SecCodeService)

    calls: list[str] = []

    async def fake_sse(self):
        calls.append("sse")
        return [SseEquityClient._row_to_official(["600000", "浦发银行", "E110    "])]

    async def fake_szse(self):
        calls.append("szse")
        return [SzseXlsxClient._row_to_official(_szse_row("000001", "平安银行"), _SZSE_HEADER_MAP)]

    monkeypatch.setattr(SecCodeService, "_fetch_sse_equity", fake_sse)
    monkeypatch.setattr(SecCodeService, "_fetch_szse_xlsx", fake_szse)

    rows = await svc._fetch_official()
    assert {r.code for r in rows} == {"600000", "000001"}
    assert sorted(calls) == ["sse", "szse"]
    assert svc._last_fetch_failed_segments == []


@pytest.mark.asyncio
async def test_fetch_official_partial_when_sse_fails(monkeypatch, capsys):
    from app.services.seccode_service import SecCodeService
    svc = SecCodeService.__new__(SecCodeService)

    async def fake_sse(self):
        raise RuntimeError("SSE down")

    async def fake_szse(self):
        return [SzseXlsxClient._row_to_official(_szse_row("000001", "平安"), _SZSE_HEADER_MAP)]

    monkeypatch.setattr(SecCodeService, "_fetch_sse_equity", fake_sse)
    monkeypatch.setattr(SecCodeService, "_fetch_szse_xlsx", fake_szse)

    rows = await svc._fetch_official()
    assert len(rows) == 1
    assert rows[0].code == "000001"
    assert svc._last_fetch_failed_segments == ["SSE_EQUITY"]
    assert "SSE" in capsys.readouterr().out


@pytest.mark.asyncio
async def test_fetch_official_partial_when_szse_fails(monkeypatch, capsys):
    from app.services.seccode_service import SecCodeService
    svc = SecCodeService.__new__(SecCodeService)

    async def fake_sse(self):
        return [SseEquityClient._row_to_official(["600000", "浦发", "E110    "])]

    async def fake_szse(self):
        raise RuntimeError("SZSE down")

    monkeypatch.setattr(SecCodeService, "_fetch_sse_equity", fake_sse)
    monkeypatch.setattr(SecCodeService, "_fetch_szse_xlsx", fake_szse)

    rows = await svc._fetch_official()
    assert len(rows) == 1
    assert rows[0].code == "600000"
    assert svc._last_fetch_failed_segments == ["SZSE_XLSX"]
    assert "SZSE" in capsys.readouterr().out


@pytest.mark.asyncio
async def test_fetch_official_both_fail_returns_empty(monkeypatch):
    from app.services.seccode_service import SecCodeService
    svc = SecCodeService.__new__(SecCodeService)

    async def fake_sse(self):
        raise RuntimeError("SSE down")

    async def fake_szse(self):
        raise RuntimeError("SZSE down")

    monkeypatch.setattr(SecCodeService, "_fetch_sse_equity", fake_sse)
    monkeypatch.setattr(SecCodeService, "_fetch_szse_xlsx", fake_szse)

    rows = await svc._fetch_official()
    assert rows == []
    assert sorted(svc._last_fetch_failed_segments) == ["SSE_EQUITY", "SZSE_XLSX"]


# --------------------------------------------------------------------- #
# sync_from_exchanges → SyncReport cookie_warmed / partial
# --------------------------------------------------------------------- #
class _FakeRepoOfficial:
    def __init__(self):
        self.entities = []
        self.saved = []

    async def find_all(self):
        return self.entities

    async def save(self, entity):
        self.saved.append(entity)
        return entity

    async def save_many(self, entities):
        for e in entities:
            await self.save(e)


@pytest.mark.asyncio
async def test_sync_from_exchanges_reports_cookie_warmed_false(monkeypatch):
    from app.services.seccode_service import SecCodeService
    svc = SecCodeService.__new__(SecCodeService)
    svc.repo = _FakeRepoOfficial()

    async def fake_ensure(self):
        raise AssertionError("_ensure_cookie should not be called")

    async def fake_fetch(self):
        return [SseEquityClient._row_to_official(["600000", "浦发", "E110    "])]

    async def fake_refill(self, codes):
        return type("R", (), {"refilled": 0, "failed": []})()

    monkeypatch.setattr(SecCodeService, "_ensure_cookie", fake_ensure, raising=False)
    monkeypatch.setattr(SecCodeService, "_fetch_official", fake_fetch)
    monkeypatch.setattr(SecCodeService, "refill_details", fake_refill)

    report = await svc.sync_from_exchanges(force=False)
    assert report.cookie_warmed is False
    assert report.partial is False


@pytest.mark.asyncio
async def test_sync_from_exchanges_partial_when_segment_fails(monkeypatch):
    from app.services.seccode_service import SecCodeService
    svc = SecCodeService.__new__(SecCodeService)
    svc.repo = _FakeRepoOfficial()

    async def fake_fetch(self):
        svc._last_fetch_failed_segments = ["SSE_EQUITY"]
        return [SseEquityClient._row_to_official(["600000", "浦发", "E110    "])]

    async def fake_refill(self, codes):
        return type("R", (), {"refilled": 0, "failed": []})()

    monkeypatch.setattr(SecCodeService, "_fetch_official", fake_fetch)
    monkeypatch.setattr(SecCodeService, "refill_details", fake_refill)

    report = await svc.sync_from_exchanges(force=False)
    assert report.partial is True
    assert report.cookie_warmed is False
